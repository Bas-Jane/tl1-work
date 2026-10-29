#!/usr/bin/env python3
"""
Excel CIQ -> TL1 generator

- Reads an Excel workbook containing one or more sheets/sections.
- Finds sections by cells that contain: "CLI Command: xxxx".
- Builds TL1 lines by matching Parameters -> CLI Name -> Target Value.
- Supports horizontal target tables and vertical target tables.
- Supports "==> Refer to <sheet> sheet" when target values are in another sheet.
- Lets you process the whole workbook, one sheet, or selected sections.
- Prompts for input workbook, output folder, and output .tl1 filename.

Required:
    pip install openpyxl
"""

from __future__ import annotations

import os
import re
import sys
from datetime import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: openpyxl is required. Install it with: pip install openpyxl")
    sys.exit(1)


# -----------------------------
# Configuration
# -----------------------------
SCRIPT_VERSION = "2.148.0"
VERSION_TAG = "v" + SCRIPT_VERSION.replace(".", "_")
MAX_TL1_LINES_PER_FILE = 10000
LARGE_SHEET_ROW_THRESHOLD = 100000
DEFAULT_CLI_SCAN_MAX_ROWS = 20000
DEFAULT_TARGET_MARKER_SCAN_MAX_ROWS = 2000
MAX_BLANK_ROWS_IN_TARGET_TABLE = 25

SKIP_VALUES = {
    "[tbd]", "[st-input]", "[telus-input]", "[enc]", "tbd", "skip", "n/a"
}
STATUS_VALUES = {"✓", "◌", "🗙", "↞", "⯑"}
TARGET_RE = re.compile(r"\btarget\s+values?\b", re.I)
CURRENT_RE = re.compile(r"\bcurrent\s+values?\b", re.I)
CLI_RE = re.compile(r"cli\s*command\s*:\s*(.+)", re.I)
REFER_RE = re.compile(r"refer\s+to\s+(.+?)\s+sheet", re.I)


@dataclass
class Section:
    sheet_name: str
    title: str
    cli_row: int
    cli_col: int
    raw_cli: str
    command: str
    cli_order: List[str]
    param_to_cli: Dict[str, str]
    key_cli_names: Optional[set[str]] = None
    param_fixed_cli_values: Optional[Dict[str, Dict[str, str]]] = None


# -----------------------------
# General helpers
# -----------------------------
def as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def norm_key(value: Any) -> str:
    return re.sub(r"\s+", " ", as_text(value)).strip().lower()


def clean_cli_name(value: Any) -> str:
    s = as_text(value)
    s = s.replace("[", "").replace("]", "")
    s = s.strip().strip(",;:")
    return s.upper()



def parse_cli_name_spec(value: Any) -> Tuple[str, Dict[str, str]]:
    """Parse a CLI Name cell and optional fixed assignments.

    Examples:
        [DSCP] ([DIR] = UL) -> ("DSCP", {"DIR": "UL"})
        [MBR]  ([DIR] = DL) -> ("MBR", {"DIR": "DL"})
        DSCP   (SVC=GTPC)   -> ("DSCP", {"SVC": "GTPC"})
        [VAL]               -> ("VAL", {})
    """
    raw = as_text(value)
    if not raw:
        return "", {}

    m = re.search(r"\[\s*([A-Za-z][A-Za-z0-9_]*)\s*\]", raw)
    if m:
        cli_name = clean_cli_name(m.group(1))
    else:
        primary = re.split(r"[\r\n(]", raw, maxsplit=1)[0]
        cli_name = clean_cli_name(primary)

    fixed: Dict[str, str] = {}

    for key, val in re.findall(
        r"\[\s*([A-Za-z][A-Za-z0-9_]*)\s*\]\s*=\s*([A-Za-z0-9_.\-]+)",
        raw,
        flags=re.I,
    ):
        fixed[clean_cli_name(key)] = as_text(val).strip().upper()

    for key, val in re.findall(
        r"(?<!\[)\b([A-Za-z][A-Za-z0-9_]*)\s*=\s*([A-Za-z0-9_.\-]+)",
        raw,
        flags=re.I,
    ):
        cli_key = clean_cli_name(key)
        if cli_key and cli_key != cli_name:
            fixed.setdefault(cli_key, as_text(val).strip().upper())

    return cli_name, fixed

def clean_param(value: Any) -> str:
    """Normalize YANG/CIQ parameter paths for reliable matching.

    CIQ templates sometimes wrap long child-list paths across lines, for example:
        amfmme-ne-grp-elem-entry/
        amfmme-ne-grp-ne-id
    while the Target Values header has the same path on one line.  Treat both as
    the same key so child parameters such as NEID/CAPA are not skipped.
    """
    s = as_text(value).strip()
    s = re.sub(r"\s*/\s*", "/", s)
    s = re.sub(r"\s+", " ", s)
    return s.strip()



def is_placeholder_param(value: Any) -> bool:
    """Return True when the Parameter/YANG cell is only a placeholder.

    CIQ sections sometimes use "-", "--", "N/A", or "None" instead of a real
    parameter name. Such values must not be used as mapping keys because several
    rows can share the same placeholder and overwrite each other.
    """
    text = clean_param(value).strip()
    normalized = text.lower()
    if not text:
        return True
    if re.fullmatch(r"-+", text):
        return True
    return normalized in {"n/a", "na", "none", "null", "[none]"}


def infer_cli_from_param(param: str, cli_order: Sequence[str]) -> Optional[str]:
    """Infer CLI name from parameter/header text when it matches the CLI list.

    This protects against CIQ rows where CLI Name is shifted/swapped and target
    headers use direct CLI names or YANG-style names.
    """
    if not param or not cli_order:
        return None

    cli_set = {x.upper() for x in cli_order}
    p_clean = clean_param(param)
    p_cli = clean_cli_name(p_clean)
    p_low = p_clean.lower()

    # Direct CLI header such as [RTALIDX] or RTALIDX.
    if p_cli in cli_set:
        return p_cli

    explicit_aliases = {
        "call-prof-index": "INDEX",
        "call-prof-vrfsp": "VRFSP",
        "call-prof-eplmn": "EPLMN",
        "call-prof-forbta": "FORBTA",
        "call-prof-forbla": "FORBLA",
        "call-prof-rqos": "RQOS",
        "call-prof-ims": "IMS",
        "call-prof-emg": "EMG",
        "call-prof-dcn_mmeg": "DCN_MMEG",
        "call-prof-nepsidx": "NEPSIDX",
        "call-prof-dcnr": "DCNR",
        "call-prof-srur": "SRUR",
        "call-prof-override": "OVERRIDE",
        "call-prof-usage": "USAGE",
        "call-prof-vplmn-addr": "VPLMN_ADDR",
        "call-prof-feat2": "FEAT2",
        "rtladix": "RTALIDX",
        "[rtalidx]": "RTALIDX",
        "call-prof-sgwrestoration": "SGWRST",
        "call-prof-nrrestriction": "NRRSTR",
        "call-prof-defaultapnindex": "DAPNIDX",
        "call-prof-ratrestirctioncause": "RATRSTRC",
        "call-prof-ratrestrictioncause": "RATRSTRC",
        "call-prof-rat-rstrc": "RATRSTRC",
        "call-prof-ratrstrc": "RATRSTRC",
        "call-prof-primarylaplmn": "PRILAPLMN",
        "call-prof-secondarylaplmn": "SNDLAPLMN",
        "call-prof-prilaplmn": "PRILAPLMN",
        "call-prof-sndlaplmn": "SNDLAPLMN",
        "call-prof-ltem": "LTEM",
    }
    if p_low in explicit_aliases and explicit_aliases[p_low] in cli_set:
        return explicit_aliases[p_low]

    suffix = p_clean.split("/")[-1].split("-")[-1].replace("-", "_").upper()
    if suffix in cli_set:
        return suffix

    p_tail = p_low.split("/")[-1]
    suffix_aliases = {
        "mtc-t3412-home": "MTC_T3412_HOME",
        "mtc-t3412-roam": "MTC_T3412_ROAM",
        "esr-ps": "ESR_PS",
        "psm-supp": "PSM_SUPP",
        "psm-t3324-pri": "PSM_T3324_PRI",
        "psm-t3412-pri": "PSM_T3412_PRI",
        "edrx-supp": "EDRX_SUPP",
        "edrx-ptw": "EDRX_PTW",
        "edrx-cycle": "EDRX_CYCLE",
        "edrx-ptw-nb": "EDRX_PTW_NB",
        "edrx-cycle-nb": "EDRX_CYCLE_NB",
        "edrx-ptw-pri": "EDRX_PTW_PRI",
        "edrx-cycle-pri": "EDRX_CYCLE_PRI",
        "dl-buff-use": "DL_BUFF_USE",
        "dl-infinite-buff": "DL_INFINITE_BUFF",
        "dl-buff-dur-offset": "DL_BUFF_DUR_OFFSET",
        "dl-buff-cnt": "DL_BUFF_CNT",
        "dl-buff-indir-use": "DL_BUFF_INDIR_USE",
        "ciot-cp": "CIOT_CP",
        "ciot-up": "CIOT_UP",
        "ciot-wopdn": "CIOT_WOPDN",
        "ciot-release-act": "CIOT_RELEASE_ACT",
        "ciot-release-time": "CIOT_RELEASE_TIME",
        "ciot-multi-tdr-rrt": "CIOT_MULTI_TDR_RRT",
        "ciot-rds-supp": "CIOT_RDS_SUPP",
        "sgwrestoration": "SGWRST",
        "nrrestriction": "NRRSTR",
        "defaultapnindex": "DAPNIDX",
        "ratrestirctioncause": "RATRSTRC",
        "ratrestrictioncause": "RATRSTRC",
        "primarylaplmn": "PRILAPLMN",
        "secondarylaplmn": "SNDLAPLMN",
        "vplmn-addr": "VPLMN_ADDR",
        "vplmn_addr": "VPLMN_ADDR",
        "dcn_mmeg": "DCN_MMEG",
        "nepsidx": "NEPSIDX",
        "ratrstrc": "RATRSTRC",
        "prilaplmn": "PRILAPLMN",
        "sndlaplmn": "SNDLAPLMN",
    }
    for key, cli in suffix_aliases.items():
        if p_tail.endswith(key) and cli in cli_set:
            return cli

    return None


def is_blank(value: Any) -> bool:
    return value is None or as_text(value) == ""


def is_skip_value(value: Any) -> bool:
    s = norm_key(value)
    return s in SKIP_VALUES or s.startswith("skip")


def is_status(value: Any) -> bool:
    return as_text(value) in STATUS_VALUES


def is_arrow_value(value: Any) -> bool:
    """Return True for CIQ ditto/carry-forward markers."""
    text = as_text(value).strip()
    return text in {"↪", "⇢", "→", "-->", "--->", "->", "=>"}


def normalize_ne_type(value: str) -> str:
    """Normalize NE-GRP NE type aliases required by TL1 output."""
    text = (value or "").upper()
    if text == "EIR-MME":
        return "DEIR"
    if text == "GMLC-MME":
        return "DGMLC"
    return text


def quote_if_needed(value: str) -> str:
    """Wrap TL1 values containing whitespace in double quotes.

    Example: NAME=Lab TOR UDM -> NAME="Lab TOR UDM".
    Existing quotes are preserved. Blank values are not quoted.
    """
    if value == "":
        return value
    if len(value) >= 2 and value[0] == '"' and value[-1] == '"':
        return value
    if any(ch.isspace() for ch in value):
        return f'"{value}"'
    return value


def normalize_rtalidx_value(value: str) -> str:
    """Normalize RTALIDX list values for MMECALL-PROF and similar commands.

    Examples:
        1 | 2       -> 1&2
        1|2         -> 1&2
        1 / 2       -> 1&2
        1, 2        -> 1&2
        [1, 2]      -> 1&2
        1 and 2     -> 1&2

    Single values are returned unchanged apart from whitespace cleanup.
    """
    text = (value or "").strip()
    if not text:
        return ""

    # Remove optional wrappers while preserving the content.
    text = text.strip().strip("[]{}()").strip()

    # Normalize common human/table separators to ampersand.
    text = re.sub(r"\s+(?:and|or)\s+", "&", text, flags=re.I)
    text = re.sub(r"\s*[|,/;]\s*", "&", text)
    text = re.sub(r"\s*&\s*", "&", text)

    # Remove extra whitespace around pure numeric tokens; keep non-numeric
    # tokens intact except for outer spacing.
    parts = [part.strip() for part in text.split("&") if part.strip()]
    if len(parts) <= 1:
        single = parts[0] if parts else ""
        if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_-]*", single or ""):
            return single.upper()
        return single

    return "&".join(parts)



def looks_like_formula_text(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")



def extract_default_value(text: str) -> Optional[str]:
    """Extract the effective value from a [default = ...] wrapper.

    Supported examples:
      [default = off]        -> off
      [default = {1}]        -> {1}
      [default = XXX [yy]]   -> XXX
      [default = abc [note]] -> abc

    Any trailing nested bracket annotation inside the wrapper is removed.
    """
    raw = as_text(text).strip()
    if not raw.lower().startswith("[default"):
        return None

    # Capture everything after '=' and before the final outer ']'.
    m = re.match(r"^\[\s*default\s*=\s*(.*)\]\s*$", raw, flags=re.I | re.S)
    if not m:
        return None

    inner = m.group(1).strip()

    # Remove trailing annotation bracket(s), e.g. "XXX [yy]" -> "XXX".
    # Preserve brace lists such as {1,2}.
    inner = re.sub(r"(?:\s+\[[^\[\]]*\])+\s*$", "", inner).strip()

    return inner



def is_reference_helper_value(value: Any) -> bool:
    """Return True for CIQ helper text that must never become TL1 data.

    Examples:
      Refer to column B below.
      Refer to column C/D below.
      See value below.
      See values below.
    """
    s = norm_key(value)
    if not s:
        return False
    return (
        s.startswith("refer to column ")
        or s in {"see value below.", "see value below", "see values below.", "see values below"}
    )


def clean_cell_value(value: Any) -> Optional[str]:
    """Return the base TL1 value text, or None when the value should be skipped/omitted."""
    if value is None:
        return ""   # keep blank as CLI_NAME=

    if is_status(value):
        return None

    if is_reference_helper_value(value):
        return None

    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = as_text(value)

    if not text:
        return ""

    # Convert [default = ...] wrappers, including nested annotations such as:
    #   [default = XXX [yy]] -> XXX
    #   [default = {1}]      -> {1}
    default_value = extract_default_value(text)
    if default_value is not None:
        text = default_value

    if is_skip_value(text):
        return None

    # Normalize complete outer-brace syntax here as well as in the final value
    # formatter. Doing it at both stages guarantees that generic renderers,
    # special renderers, and values unwrapped from [default = ...] all behave
    # consistently.
    text = normalize_braced_tl1_value(text)

    # Remove surrounding quotes only; field-specific formatting is applied later.
    text = text.strip().strip('"').strip("'")
    return text



def normalize_braced_tl1_value(value: str) -> str:
    """Convert CIQ brace notation into TL1 value notation.

    Examples:
      {XY}          -> XY
      {XY, ZX}      -> XY&ZX
      {XY,ZX,AB}    -> XY&ZX&AB

    The conversion is applied only when the complete trimmed value is enclosed
    by one outer pair of braces. Empty list items are ignored.
    """
    raw = as_text(value).strip()
    if len(raw) < 2 or not (raw.startswith("{") and raw.endswith("}")):
        return raw

    inner = raw[1:-1].strip()
    if not inner:
        return ""

    parts = [part.strip() for part in inner.split(",")]
    parts = [part for part in parts if part]
    return "&".join(parts)


def format_tl1_value(cli_name: str, value: str, command: str = "") -> str:
    """Apply TL1 output case rules for a specific CLI parameter.

    Most TL1 values are uppercased, but values assigned to POD must be kept
    lowercase because Kubernetes/OpenShift pod/service interface names are
    case-sensitive in the generated command.
    """
    value = normalize_braced_tl1_value(value)
    if value == "":
        return ""

    cli = (cli_name or "").strip().upper()
    cmd = (command or "").strip().upper()

    # CHG-FD-LMT value normalization.
    # The service/container name "istio-proxy" is case-sensitive and must remain
    # lowercase instead of being converted to ISTIO-PROXY.
    if cmd == "CHG-FD-LMT" and value.strip().lower() == "istio-proxy":
        return "istio-proxy"

    # CHG-DIAG-INF MSG normalization.
    # These source values are YANG-style strings but TL1 expects short codes:
    #   status-inhibit -> INH
    #   status-allow   -> ALW
    if cmd == "CHG-DIAG-INF" and cli == "MSG":
        normalized = value.strip().lower()
        if normalized == "status-inhibit":
            return "INH"
        if normalized == "status-allow":
            return "ALW"

    # RTALIDX can be entered as "1 | 2", "1|2", "1 / 2", "1, 2", etc.
    # TL1 expects ampersand-separated values, especially for CHG-MMECALL-PROF:
    #   RTALIDX=1&2
    if cli == "RTALIDX":
        return quote_if_needed(normalize_rtalidx_value(value))

    # Keep POD values lowercase globally.
    # Example: POD=SVC-MME-ANINTF-SCTP-0 -> POD=svc-mme-anintf-sctp-0
    if cli == "POD":
        return quote_if_needed(value.lower())

    # DNN values must remain lowercase.
    if cli == "DNN":
        return quote_if_needed(value.lower())

    # CHG-DSCP-INF requires SVC values in uppercase.
    # Example: SVC=gtpc -> SVC=GTPC
    if cmd == "CHG-DSCP-INF" and cli == "SVC":
        return quote_if_needed(value.upper())

    # CRTE-RAD-INTF case rules:
    #   AF must be uppercase (IP4 / IP6)
    #   KEY must remain lowercase
    if cmd == "CRTE-RAD-INTF" and cli == "AF":
        return quote_if_needed(value.upper())
    if cmd == "CRTE-RAD-INTF" and cli == "KEY":
        return quote_if_needed(value.lower())

    # CRTE-RACCT requires KEY values in lowercase.
    # Example: KEY=ABC-XYZ -> KEY=abc-xyz
    if cmd == "CRTE-RACCT" and cli == "KEY":
        return quote_if_needed(value.lower())

    # CHG-NFM-CONF SUBSNFTYPE accepts an ampersand-separated list.
    # Example: {AMF, UDM} -> AMF&UDM
    if cmd == "CHG-NFM-CONF" and cli == "SUBSNFTYPE":
        raw = value.strip()
        if raw.startswith("{") and raw.endswith("}"):
            raw = raw[1:-1].strip()
        parts = [p.strip() for p in re.split(r"[,;|&]+", raw) if p.strip()]
        normalized = "&".join(p.upper() for p in parts)
        return quote_if_needed(normalized)

    # Keep prior important case-sensitive fields stable when this simplified
    # script is used directly.
    preserve_case = {
        "NFID", "FQDN", "SVCID", "INSTID", "PGWFQDN", "DFTDNN", "APN",
        "NAME", "DESC",
    }
    lower_case = {"SVC", "HOST", "REALM"}
    if cli in preserve_case:
        return quote_if_needed(value)
    if cli in lower_case:
        return quote_if_needed(value.lower())
    # CRTE-SVC-VER requires NF values in uppercase.
    # Example: NF=smf -> NF=SMF
    if cmd == "CRTE-SVC-VER" and cli == "NF":
        return quote_if_needed(value.upper())

    text = value.upper()
    if cli == "TYPE" and text == "MEMORY":
        text = "MEM"
    return quote_if_needed(text)


def iter_used_cells(ws) -> Iterable[Tuple[int, int, Any]]:
    for row in ws.iter_rows(min_col=1, max_col=scan_column_limit(ws)):
        for cell in row:
            if cell.value is not None:
                yield cell.row, cell.column, cell.value


def get_effective_value(ws_values, row: int, col: int) -> Any:
    return ws_values.cell(row=row, column=col).value


def row_values(ws, row: int, start_col: int = 1, end_col: Optional[int] = None) -> List[Any]:
    if end_col is None:
        end_col = scan_column_limit(ws)
    return [ws.cell(row=row, column=c).value for c in range(start_col, end_col + 1)]


# -----------------------------
# CLI and section parsing
# -----------------------------

def extract_fixed_cli_values(raw_cli_text: str) -> Dict[str, str]:
    """Extract fixed KEY=VALUE assignments from a CLI Command template."""
    m = CLI_RE.search(raw_cli_text)
    cli = m.group(1).strip() if m else raw_cli_text.strip()

    fixed: Dict[str, str] = {}
    command_part, sep, suffix = cli.partition(";")

    if ":" in command_part:
        command_name, rest = command_part.split(":", 1)
        for part in rest.split(","):
            token = part.strip().strip("[]")
            if "=" not in token:
                continue
            key, value = token.split("=", 1)
            key = clean_cli_name(key)
            value = as_text(value).strip().strip('"').strip("'")
            if key and value:
                fixed[key] = format_tl1_value(key, value, command_name.strip())

    if sep and suffix:
        command_name = command_part.split(":", 1)[0].strip()
        for key, value in re.findall(
            r"\b([A-Za-z][A-Za-z0-9_]*)\s*=\s*([A-Za-z0-9_.\-']+)",
            suffix,
        ):
            cli_key = clean_cli_name(key)
            if cli_key:
                fixed[cli_key] = format_tl1_value(cli_key, value, command_name)

    return fixed


def parse_cli(raw_cli_text: str) -> Tuple[str, List[str]]:
    """Parse a CLI template into command name and ordered CLI field names.

    Fixed assignments such as TYPE=PAGE are represented in cli_order as TYPE.
    """
    m = CLI_RE.search(raw_cli_text)
    cli = m.group(1).strip() if m else raw_cli_text.strip()
    cli = cli.rstrip()

    command_part = cli.split(";", 1)[0].strip()
    if ":" not in command_part:
        return command_part.rstrip(";").upper(), []

    command, rest = command_part.split(":", 1)
    parts: List[str] = []
    for part in rest.split(","):
        token = part.strip().strip("[]")
        if not token:
            continue
        if "=" in token:
            token = token.split("=", 1)[0].strip()
        name = clean_cli_name(token)
        if name and name not in parts:
            parts.append(name)

    return command.strip().upper(), parts

def find_section_title(ws_values, cli_row: int) -> str:
    """Find nearest non-empty title above CLI command. Preferred title is a cell with '(...)'."""
    best = ""
    for r in range(cli_row - 1, max(0, cli_row - 12), -1):
        for c in range(1, scan_column_limit(ws_values) + 1):
            text = as_text(ws_values.cell(r, c).value)
            if not text:
                continue
            low = text.lower()
            if any(x in low for x in ["yang path", "review status", "config status"]):
                continue
            if "(" in text and ")" in text:
                return text
            if not best and len(text) > 2:
                best = text
    return best or ws_values.title



def is_out_of_scope_section(ws_values, section: Section, next_cli_row: int) -> bool:
    """Return True when THIS CLI section is marked Out of Scope.

    Review Status / Config Status are located above the section's CLI Command.
    Scan only the current section header area ending at section.cli_row.  Do not
    scan forward toward the next CLI command because the next section's header
    may start several rows before its CLI row and could otherwise cause a false
    Out-of-Scope match for the current section.
    """
    start_row = max(1, section.cli_row - 15)
    end_row = section.cli_row

    for r in range(start_row, end_row + 1):
        row_values_text = [
            as_text(ws_values.cell(r, c).value)
            for c in range(1, scan_column_limit(ws_values) + 1)
        ]
        row_norm = [re.sub(r"\s+", " ", v).strip().lower() for v in row_values_text if v]

        if not row_norm:
            continue

        has_status_label = any(
            v in {"review status", "config status"} or
            v.startswith("review status") or
            v.startswith("config status")
            for v in row_norm
        )
        has_out_of_scope = any(
            v == "out of scope" or "out of scope" in v
            for v in row_norm
        )

        if has_status_label and has_out_of_scope:
            return True

    return False


def is_default_wrapped_value(value: Any) -> bool:
    """True only for a complete CIQ default wrapper such as [default = 24000]."""
    if value is None:
        return False
    text = as_text(value).strip()
    return bool(re.fullmatch(r"\[\s*default\s*=\s*[^\]]+\]", text, flags=re.I))


def find_key_cli_names(
    ws_values,
    cli_row: int,
    next_cli_row: int,
) -> set[str]:
    """Return CLI names whose parameter Type is a key leaf.

    Supports values such as:
        key leaf
        child list key leaf
        child-list key leaf
    """
    max_row = min(ws_values.max_row, next_cli_row - 1 if next_cli_row else ws_values.max_row)
    header_row = None

    for r in range(cli_row, max_row + 1):
        texts = [norm_key(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1)]
        if any(t in {"parameters", "paramaters"} for t in texts):
            header_row = r
            break

    if header_row is None:
        return set()

    type_col = None
    cli_col = None
    for c in range(1, scan_column_limit(ws_values) + 1):
        key = norm_key(ws_values.cell(header_row, c).value)
        if key == "type":
            type_col = c
        elif key == "cli name":
            cli_col = c

    if type_col is None or cli_col is None:
        return set()

    keys: set[str] = set()
    for r in range(header_row + 1, max_row + 1):
        row_text = " | ".join(
            as_text(ws_values.cell(r, c).value)
            for c in range(1, scan_column_limit(ws_values) + 1)
        )
        if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
            break

        type_text = norm_key(ws_values.cell(r, type_col).value)
        cli_name = clean_cli_name(ws_values.cell(r, cli_col).value)
        if cli_name and "key leaf" in type_text:
            keys.add(cli_name)

    return keys


def record_has_only_default_non_key_values(section: Section, record: Dict[str, Any]) -> bool:
    """Return True when every assigned non-key parameter is [default = ...].

    Key fields identify the configuration instance and do not prevent skipping.
    Blank/status/arrow/control fields are ignored. At least one non-key assigned
    parameter must exist, and every such parameter must be a complete default
    wrapper.

    This is generic across all commands.
    """
    key_cli_names = {x.upper() for x in (section.key_cli_names or set())}
    non_key_raw_values: List[Any] = []

    for param, raw_value in record.items():
        if str(param).startswith("__"):
            continue

        cli_name = section.param_to_cli.get(param)
        if not cli_name:
            # Direct CLI-name records used by special parsers.
            candidate = clean_cli_name(param)
            if candidate in {x.upper() for x in section.cli_order}:
                cli_name = candidate

        if not cli_name:
            continue

        cli_upper = cli_name.upper()
        if cli_upper in key_cli_names:
            continue

        if raw_value is None or is_blank(raw_value) or is_status(raw_value) or is_arrow_value(raw_value):
            continue

        non_key_raw_values.append(raw_value)

    return bool(non_key_raw_values) and all(
        is_default_wrapped_value(v) for v in non_key_raw_values
    )



def find_parameter_fixed_cli_values(
    ws_values,
    cli_row: int,
    next_cli_row: int,
) -> Dict[str, Dict[str, str]]:
    """Return fixed CLI assignments attached to parameter rows.

    This captures CIQ annotations such as:
        [DSCP] ([DIR] = UL)
        [DSCP] ([DIR] = DL)
    """
    max_row = min(ws_values.max_row, next_cli_row - 1 if next_cli_row else ws_values.max_row)
    header_row = None
    for r in range(cli_row, max_row + 1):
        texts = [norm_key(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1)]
        if any(t in {"parameters", "paramaters"} for t in texts):
            header_row = r
            break
    if header_row is None:
        return {}

    param_col = None
    cli_col = None
    for c in range(1, scan_column_limit(ws_values) + 1):
        key = norm_key(ws_values.cell(header_row, c).value)
        if key in {"parameters", "paramaters"}:
            param_col = c
        elif key == "cli name":
            cli_col = c
    if param_col is None or cli_col is None:
        return {}

    result: Dict[str, Dict[str, str]] = {}
    for r in range(header_row + 1, max_row + 1):
        row_text = " | ".join(
            as_text(ws_values.cell(r, c).value)
            for c in range(1, scan_column_limit(ws_values) + 1)
        )
        if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
            break

        param = clean_param(ws_values.cell(r, param_col).value)
        if is_placeholder_param(param):
            continue
        _cli, fixed = parse_cli_name_spec(ws_values.cell(r, cli_col).value)
        if fixed:
            result[param] = fixed
    return result


def find_parameter_mapping(ws_values, cli_row: int, cli_col: int, next_cli_row: int) -> Dict[str, str]:
    """Find Parameters table after CLI row and return parameter leaf -> CLI name."""
    max_row = min(ws_values.max_row, next_cli_row - 1 if next_cli_row else ws_values.max_row)
    header_row = None

    for r in range(cli_row, max_row + 1):
        texts = [norm_key(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1)]
        if any(t in {"parameters", "paramaters"} for t in texts):
            header_row = r
            break

    if header_row is None:
        return {}

    # Locate PARAMATERS and CLI Name columns.
    param_col = None
    cli_name_col = None
    for c in range(1, scan_column_limit(ws_values) + 1):
        text = norm_key(ws_values.cell(header_row, c).value)
        if text in {"parameters", "paramaters"}:
            param_col = c
        elif text == "cli name":
            cli_name_col = c

    if param_col is None or cli_name_col is None:
        return {}

    mapping: Dict[str, str] = {}
    for r in range(header_row + 1, max_row + 1):
        # Stop at first Target Values block, next section title, or next CLI command.
        row_text = " | ".join(as_text(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1))
        if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
            break

        param = clean_param(ws_values.cell(r, param_col).value)
        cli_name, _fixed_values = parse_cli_name_spec(ws_values.cell(r, cli_name_col).value)

        # Placeholder Parameter values such as "-" are not real parameter
        # names. Do not use them as mapping keys or later rows will overwrite
        # earlier rows that use the same placeholder.
        if is_placeholder_param(param):
            param = ""

        # Normal mapping: YANG/CIQ parameter path -> CLI Name.
        if param and cli_name:
            mapping[param] = cli_name

        # Fallback mapping: CLI Name -> CLI Name.
        # Some sheets, including MMERTAL-INF, keep the Parameter/YANG path cell
        # blank and use CLI headers directly in the Target Values table:
        #   RTALIDX | CAUSE | MCC | MNC | TAC
        # Without this alias the parser cannot match the target headers and
        # emits "No TL1 rows generated".
        if cli_name:
            mapping.setdefault(cli_name, cli_name)

    return mapping



# -----------------------------
# Performance helpers
# -----------------------------
DEFAULT_SCAN_MAX_COLUMNS = 256

def scan_column_limit(ws_values) -> int:
    """Return a practical upper column bound for CIQ scanning.

    Excel files often report a very large max_column because of formatting,
    merged cells, or stale styles far to the right of the real CIQ data.
    Scanning all of those columns repeatedly is extremely expensive.

    Override with:
        TL1_SCAN_MAX_COLUMNS=0   -> scan all worksheet columns
        TL1_SCAN_MAX_COLUMNS=N   -> scan at most N columns

    The default 256 columns comfortably includes the normal CIQ area,
    including requirement/value columns such as AA and beyond.
    """
    raw = os.environ.get("TL1_SCAN_MAX_COLUMNS", str(DEFAULT_SCAN_MAX_COLUMNS)).strip()
    try:
        configured = int(raw)
    except ValueError:
        configured = DEFAULT_SCAN_MAX_COLUMNS

    if configured <= 0:
        return ws_values.max_column
    return min(ws_values.max_column, max(32, configured))


# Cache case-insensitive parameter lookup tables by the identity of the
# known_params set. The cache retains the set object, so Python cannot recycle
# its id for a different set while the cached entry exists.
_PARAM_MATCH_CACHE: Dict[int, Tuple[set[str], Dict[str, str]]] = {}

def _known_param_lower_map(known_params: set[str]) -> Dict[str, str]:
    key = id(known_params)
    cached = _PARAM_MATCH_CACHE.get(key)
    if cached is not None and cached[0] is known_params:
        return cached[1]

    lower_map = {str(k).lower(): k for k in known_params}
    _PARAM_MATCH_CACHE[key] = (known_params, lower_map)

    # Keep memory bounded for workbooks containing many sections.
    if len(_PARAM_MATCH_CACHE) > 512:
        # Dropping the oldest-ish entries is sufficient; exact LRU behavior is
        # unnecessary here.
        for old_key in list(_PARAM_MATCH_CACHE.keys())[:256]:
            _PARAM_MATCH_CACHE.pop(old_key, None)

    return lower_map


def discover_sections(wb_values, sheets_to_scan: Sequence[str]) -> List[Section]:
    sections: List[Section] = []
    for sheet_name in sheets_to_scan:
        ws = wb_values[sheet_name]
        cli_cells = []

        # Large data sheets can contain hundreds of thousands of rows. CLI Command
        # sections are normally in the CIQ section/header area, so scan only the
        # first N rows by default to avoid a full 400k+ row pass on data-only sheets.
        # Set TL1_CLI_SCAN_MAX_ROWS=0 to force a full-sheet CLI scan.
        scan_limit_env = os.environ.get("TL1_CLI_SCAN_MAX_ROWS", str(DEFAULT_CLI_SCAN_MAX_ROWS)).strip()
        try:
            scan_limit = int(scan_limit_env)
        except ValueError:
            scan_limit = DEFAULT_CLI_SCAN_MAX_ROWS
        max_scan_row = ws.max_row if scan_limit <= 0 else min(ws.max_row, scan_limit)
        if ws.max_row > max_scan_row:
            print(f"Scanning first {max_scan_row:,} rows for CLI sections on large sheet '{sheet_name}' "
                  f"({ws.max_row:,} total rows). Set TL1_CLI_SCAN_MAX_ROWS=0 to scan all rows.")

        for row in ws.iter_rows(
            min_row=1,
            max_row=max_scan_row,
            min_col=1,
            max_col=scan_column_limit(ws),
        ):
            for cell in row:
                if cell.value is None:
                    continue
                text = as_text(cell.value)
                if CLI_RE.search(text):
                    cli_cells.append((cell.row, cell.column, text))
        cli_cells.sort()

        for i, (r, c, raw) in enumerate(cli_cells):
            next_cli_row = cli_cells[i + 1][0] if i + 1 < len(cli_cells) else ws.max_row + 1
            command, cli_order = parse_cli(raw)
            title = find_section_title(ws, r)
            mapping = find_parameter_mapping(ws, r, c, next_cli_row)
            key_cli_names = find_key_cli_names(ws, r, next_cli_row)
            param_fixed_cli_values = find_parameter_fixed_cli_values(ws, r, next_cli_row)
            # Keep the explicit CLI Name column as the source of truth.
            # Earlier versions inferred CLI names from the YANG leaf suffix for every row,
            # but that broke valid child-list mappings such as:
            #   smf-inf-smf-svc-entry/smf-inf-smf-svc-name -> SVC
            # because the suffix "name" was incorrectly remapped to NAME.
            # Only infer when the explicit CLI name is missing or not part of the command.
            for _param, _cli in list(mapping.items()):
                if _cli and _cli in cli_order:
                    continue
                _inferred = infer_cli_from_param(_param, cli_order)
                if _inferred:
                    mapping[_param] = _inferred

            # Add direct aliases for every CLI placeholder. This makes Target
            # Values headers like INDEX, [RTALIDX], or SGWRST parse even when
            # the Parameter/YANG-path column is blank or shifted.
            for _cli in cli_order:
                mapping.setdefault(_cli.upper(), _cli.upper())
                mapping.setdefault(f"[{_cli.upper()}]", _cli.upper())
            sections.append(
                Section(
                    sheet_name=sheet_name,
                    title=title,
                    cli_row=r,
                    cli_col=c,
                    raw_cli=raw,
                    command=command,
                    cli_order=cli_order,
                    param_to_cli=mapping,
                    key_cli_names=key_cli_names,
                    param_fixed_cli_values=param_fixed_cli_values,
                )
            )
    return sections




def find_parameter_table_end(ws_values, section: Section, next_cli_row: int) -> int:
    """Return the last row in the Parameters table for this section.

    The Parameters header row may itself contain a "Target Value" helper column,
    so we must first find the PARAMETERS/PARAMATERS row, then stop at the next
    true Target Values block below the parameter rows.
    """
    known_params = set(section.param_to_cli.keys())
    if not known_params:
        return section.cli_row

    max_row = min(ws_values.max_row, next_cli_row - 1)
    header_row = None
    for r in range(section.cli_row + 1, max_row + 1):
        row_texts = [norm_key(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1)]
        if any(t in {"parameters", "paramaters"} for t in row_texts):
            header_row = r
            break
    if header_row is None:
        return section.cli_row

    last = header_row
    for r in range(header_row + 1, max_row + 1):
        row_join = " | ".join(as_text(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1))
        if TARGET_RE.search(row_join):
            break
        for c in range(1, scan_column_limit(ws_values) + 1):
            if clean_param(ws_values.cell(r, c).value) in known_params:
                last = r
                break
    return last

# -----------------------------
# Target table discovery and parsing
# -----------------------------

def is_real_target_marker_row(ws_values, row: int, col: int) -> bool:
    """Return True for a real Target Values block marker.

    Some damaged CIQ rows can contain the literal text "Target Values (B1C3)"
    inside what is actually a data row, for example:
        Target Values (B1C3) | [ENC] | ... | Target Values (B1C3) | [ENC]

    A real marker may contain harmless notes such as "Entries may be hidden",
    but should not contain an actual target value immediately to the right.
    """
    marker_text = as_text(ws_values.cell(row, col).value)
    if not TARGET_RE.search(marker_text):
        return False

    harmless = {
        "",
        "entries may be hidden",
        "entry may be hidden",
        "values may be hidden",
        "values hidden",
        "value hidden",
    }

    # Inspect only a small target-side window before Current Values.
    for c in range(col + 1, min(scan_column_limit(ws_values), col + 4) + 1):
        v = ws_values.cell(row, c).value
        if is_blank(v) or is_status(v):
            continue
        s = norm_key(v)
        if s in harmless:
            continue
        if CURRENT_RE.search(as_text(v)):
            break

        # Any other value means this is likely a damaged data row, not a marker.
        return False

    return True


def find_target_markers(ws_values, start_row: int = 1, end_row: Optional[int] = None) -> List[Tuple[int, int]]:
    """Find Target Values markers without scanning huge data sheets end-to-end.

    When end_row is not supplied, this is usually a referred data sheet.  In CIQ
    workbooks the Target Values marker is normally near the top, while data may
    extend hundreds of thousands of rows.  To avoid scanning 400k rows just to
    locate the marker, the default search is capped by TL1_TARGET_MARKER_SCAN_MAX_ROWS.
    Set TL1_TARGET_MARKER_SCAN_MAX_ROWS=0 to force a full scan.
    """
    if end_row is None:
        scan_limit_env = os.environ.get("TL1_TARGET_MARKER_SCAN_MAX_ROWS", str(DEFAULT_TARGET_MARKER_SCAN_MAX_ROWS)).strip()
        try:
            scan_limit = int(scan_limit_env)
        except ValueError:
            scan_limit = DEFAULT_TARGET_MARKER_SCAN_MAX_ROWS
        if scan_limit <= 0:
            end_row = ws_values.max_row
        else:
            end_row = min(ws_values.max_row, start_row + scan_limit - 1)
            if ws_values.max_row > end_row:
                print(
                    f"Scanning first {end_row:,} rows for Target Values on large/referred sheet "
                    f"'{ws_values.title}' ({ws_values.max_row:,} total rows). "
                    "Set TL1_TARGET_MARKER_SCAN_MAX_ROWS=0 to scan all rows."
                )

    markers = []
    max_r = min(end_row, ws_values.max_row)
    for row in ws_values.iter_rows(
        min_row=start_row,
        max_row=max_r,
        min_col=1,
        max_col=scan_column_limit(ws_values),
    ):
        for cell in row:
            if cell.value is None:
                continue
            text = as_text(cell.value)
            if TARGET_RE.search(text) and is_real_target_marker_row(
                ws_values, cell.row, cell.column
            ):
                markers.append((cell.row, cell.column))
    return markers

def find_referred_sheet_name(ws_values, start_row: int, end_row: Optional[int] = None) -> Optional[str]:
    if end_row is None:
        end_row = ws_values.max_row
    for r in range(start_row, min(end_row, ws_values.max_row) + 1):
        for c in range(1, scan_column_limit(ws_values) + 1):
            text = as_text(ws_values.cell(r, c).value)
            m = REFER_RE.search(text)
            if m:
                return m.group(1).strip()
    return None


def resolve_referred_sheet(wb_values, referred: str) -> Optional[str]:
    if not referred:
        return None
    referred_norm = norm_key(referred)
    for name in wb_values.sheetnames:
        if norm_key(name) == referred_norm:
            return name
    # Loose match: ignore punctuation/spaces.
    compact = re.sub(r"[^a-z0-9]", "", referred_norm)
    for name in wb_values.sheetnames:
        if re.sub(r"[^a-z0-9]", "", norm_key(name)) == compact:
            return name
    return None


def _param_match(value: Any, known_params: set[str]) -> Optional[str]:
    text = clean_param(value)
    if text in known_params:
        return text

    # Accept direct CLI-name headers in target tables, including optional brackets.
    # Example: [RTALIDX] should match the RTALIDX -> RTALIDX alias in known_params.
    cli_text = clean_cli_name(text)
    if cli_text in known_params:
        return cli_text

    # Accept case-insensitive parameter/header match. The lookup map is cached
    # per known_params set; older versions rebuilt this dictionary for every
    # cell examined, which became a major CPU hot spot on large sheets.
    known_lower = _known_param_lower_map(known_params)
    text_lower = text.lower()
    cli_lower = cli_text.lower()
    if text_lower in known_lower:
        return known_lower[text_lower]
    if cli_lower in known_lower:
        return known_lower[cli_lower]

    return None


def _region_end_row(ws_values, marker_row: int, next_section_row: Optional[int] = None) -> int:
    """Return the last row belonging to one Target Values block.

    A CIQ section may contain many Target Values blocks. Older versions let
    special sequential parsers run until the next CLI section, which caused
    later objects to overwrite/duplicate earlier objects.
    """
    limit = min(
        ws_values.max_row,
        (next_section_row - 1) if next_section_row else ws_values.max_row,
    )

    for r in range(marker_row + 1, limit + 1):
        # A new CLI section always ends this region.
        for c in range(1, min(scan_column_limit(ws_values), 16) + 1):
            if CLI_RE.search(as_text(ws_values.cell(r, c).value)):
                return r - 1

        # Stop only on a *real* Target Values marker, not on damaged data rows.
        for c in range(1, min(scan_column_limit(ws_values), 16) + 1):
            if is_real_target_marker_row(ws_values, r, c):
                return r - 1

    return limit




def target_column_bounds(ws_values, marker_row: int, marker_col: int) -> Tuple[int, int]:
    """Return inclusive target value column bounds, stopping before Current Values."""
    end_col = scan_column_limit(ws_values)
    for c in range(marker_col + 1, scan_column_limit(ws_values) + 1):
        if CURRENT_RE.search(as_text(ws_values.cell(marker_row, c).value)):
            end_col = c - 1
            break
    return marker_col, end_col


def normalize_row_action(value: Any) -> Optional[str]:
    """Return a normalized row action marker such as CHG or CRTE.

    CIQ target tables sometimes have an extra leading column whose cell contains
    CHG on rows that must use a change command instead of the section's normal
    create command. The action column may have a blank/unmapped header.
    """
    text = as_text(value).strip().upper()
    if text in {"CHG", "CHANGE"}:
        return "CHG"
    if text in {"CRTE", "CREATE"}:
        return "CRTE"
    if text in {"REG", "REGISTER"}:
        return "REG"
    return None


def derive_command_for_row(base_command: str, action: Optional[str]) -> str:
    """Derive the TL1 command prefix for a row action marker.

    Examples:
        CRTE-FOO + CHG  -> CHG-FOO
        CRTE-FOO + CRTE -> CRTE-FOO
        REG-FOO  + CHG  -> CHG-FOO

    If no supported action marker is present, return the section command.
    """
    cmd = (base_command or "").strip().upper()
    if not action:
        return cmd

    action = action.upper()
    if "-" in cmd:
        _prefix, suffix = cmd.split("-", 1)
    else:
        suffix = cmd

    if action == "CHG":
        return f"CHG-{suffix}"
    if action == "CRTE":
        return f"CRTE-{suffix}"
    if action == "REG":
        return f"REG-{suffix}"
    return cmd


def detect_row_action_near_target_columns(
    ws_values,
    row: int,
    param_cols: Dict[int, str],
    marker_col: int,
) -> Optional[str]:
    """Detect CHG/CRTE/REG markers in nearby unmapped cells on a target row.

    The common layout has a blank first/action column immediately to the left
    of the first mapped parameter column, but this scans the small gap between
    the Target Values marker and the first mapped parameter as well.
    """
    if not param_cols:
        return None

    first_param_col = min(param_cols.keys())

    # Prefer the immediately preceding cell, then scan left back to marker_col.
    candidate_cols = []
    if first_param_col > 1:
        candidate_cols.append(first_param_col - 1)

    left_bound = max(1, min(marker_col, first_param_col - 1))
    for c in range(first_param_col - 2, left_bound - 1, -1):
        if c not in candidate_cols:
            candidate_cols.append(c)

    # Also inspect a few leading columns because some sheets place the action
    # column outside the merged Target Values area.
    for c in range(1, min(first_param_col, 5)):
        if c not in candidate_cols:
            candidate_cols.append(c)

    for c in candidate_cols:
        action = normalize_row_action(ws_values.cell(row, c).value)
        if action:
            return action

    return None



def row_starts_documentation_block(ws_values, row: int, scan_cols: Optional[Sequence[int]] = None) -> bool:
    """Detect detail/reference tables that follow a real Target Values table.

    CIQ sheets commonly append documentation blocks after configuration rows:
      Resource Details | Details hidden
      Call Timer Details | ...
      License Details | ...
      Resources (...) | Description

    These are not configuration records and must terminate target parsing.
    """
    if scan_cols is None:
        scan_cols = range(1, scan_column_limit(ws_values) + 1)

    vals = [
        re.sub(r"\s+", " ", as_text(ws_values.cell(row, c).value)).strip().lower()
        for c in scan_cols
        if as_text(ws_values.cell(row, c).value)
    ]
    if not vals:
        return False

    joined = " | ".join(vals)

    strong_markers = (
        "resource details",
        "details hidden",
        "call timer details",
        "license details",
        "timer details",
    )
    if any(marker in joined for marker in strong_markers):
        return True

    # Generic reference/detail-table header after target rows.
    if any(v == "description" for v in vals):
        if any(
            ("resource" in v or "timer" in v or "item" == v or "value range" in v)
            for v in vals
        ):
            return True

    return False



def row_starts_next_section_header(
    ws_values,
    row: int,
    scan_cols: Optional[Sequence[int]] = None,
) -> bool:
    """Detect the header/status rows of the following CIQ section.

    A new section title often appears several rows before its CLI Command, so
    limiting parsing only by the next CLI row is not enough. Review Status and
    Config Status rows after target records have begun mark the end of the
    current target table.
    """
    if scan_cols is None:
        scan_cols = range(1, scan_column_limit(ws_values) + 1)

    vals = [
        re.sub(r"\s+", " ", as_text(ws_values.cell(row, c).value)).strip().lower()
        for c in scan_cols
        if as_text(ws_values.cell(row, c).value)
    ]
    if not vals:
        return False

    joined = " | ".join(vals)
    return (
        "review status" in joined
        or "config status" in joined
        or any(v.startswith("yang path:") for v in vals)
        or any(v.startswith("cli command:") for v in vals)
    )



def parse_horizontal_targets(
    ws_values,
    marker_row: int,
    marker_col: int,
    known_params: set[str],
    max_end_row: Optional[int] = None,
) -> List[Dict[str, Any]]:
    """Parse horizontal Target Values tables efficiently, including very large tables.

    The older loop scanned every cell in every row to look for stop markers. On
    400k-row target sheets this made runtime explode.  This version only reads
    the target columns plus a small leading window used to detect a new section.
    """
    end_row = _region_end_row(ws_values, marker_row, max_end_row)
    start_col, end_col = target_column_bounds(ws_values, marker_row, marker_col)

    header_row = None
    param_cols: Dict[int, str] = {}

    # Header is normally 1-3 rows below Target Values marker.
    for r in range(marker_row + 1, min(end_row, marker_row + 5) + 1):
        matches: Dict[int, str] = {}
        for c in range(start_col, end_col + 1):
            p = _param_match(ws_values.cell(r, c).value, known_params)
            if p:
                matches[c] = p
        if len(matches) >= 2:
            header_row = r
            param_cols = matches
            break

    if header_row is None:
        return []

    records: List[Dict[str, Any]] = []
    blank_run = 0
    inherited_values: Dict[str, Any] = {}
    watch_cols = sorted(set(list(param_cols.keys()) + list(range(1, min(ws_values.max_column, 8) + 1))))

    for r in range(header_row + 1, end_row + 1):
        # Only inspect a small subset of columns for section/current markers.
        marker_text = " | ".join(as_text(ws_values.cell(r, c).value) for c in watch_cols)
        if CLI_RE.search(marker_text):
            break
        if TARGET_RE.search(marker_text) and r != marker_row:
            break
        if CURRENT_RE.search(marker_text):
            break

        # Stop before post-target documentation/reference tables or the next
        # CIQ section header, which may appear before the next CLI Command row.
        if records and (
            row_starts_documentation_block(ws_values, r, watch_cols)
            or row_starts_next_section_header(ws_values, r, watch_cols)
        ):
            break

        rec: Dict[str, Any] = {}
        has_any_value = False
        for c, p in param_cols.items():
            value = ws_values.cell(r, c).value
            if value is not None and as_text(value) != "":
                has_any_value = True
            # Horizontal continuation arrows inherit the previous value for
            # that parameter. This is used by PCF-ERR and similar tables.
            if is_arrow_value(value):
                value = inherited_values.get(p, value)
            elif value is not None and as_text(value) != "":
                inherited_values[p] = value

            rec[p] = value

        # Capture an optional row-level action marker from an unmapped control
        # column. This allows a row containing CHG to render as CHG-XXXX even
        # when the section CLI Command is CRTE-XXXX.
        row_action = detect_row_action_near_target_columns(
            ws_values, r, param_cols, marker_col
        )
        if row_action:
            rec["__ROW_ACTION__"] = row_action

        if has_any_value:
            if record_contains_metadata_row(rec):
                # Once real target rows have started, a metadata/details row
                # marks the end of the target table. Stop parsing so following
                # documentation rows (Resource Details, descriptions, timer
                # details, etc.) cannot become TL1 commands.
                if records:
                    break
                continue
            blank_run = 0
            records.append(rec)
            if len(records) % 50000 == 0:
                print(f"  Parsed {len(records):,} target rows from sheet '{ws_values.title}'...")
        else:
            blank_run += 1
            # On very large tables there may be occasional blank separators, but
            # repeated blanks should terminate the table.
            if blank_run >= MAX_BLANK_ROWS_IN_TARGET_TABLE:
                break

    return records

def parse_vertical_targets(
    ws_values,
    marker_row: int,
    marker_col: int,
    known_params: set[str],
    max_end_row: Optional[int] = None,
) -> List[Dict[str, Any]]:
    """Parse target table where parameter names are vertical and values are in column(s) to the right."""
    end_row = _region_end_row(ws_values, marker_row, max_end_row)
    start_col, end_col = target_column_bounds(ws_values, marker_row, marker_col)

    # Find dominant parameter column after marker.
    best_col = None
    best_count = 0
    scan_end = min(end_row, marker_row + 100)
    for c in range(start_col, end_col + 1):
        count = 0
        for r in range(marker_row + 1, scan_end + 1):
            if _param_match(ws_values.cell(r, c).value, known_params):
                count += 1
        if count > best_count:
            best_count = count
            best_col = c

    if best_col is None or best_count < 2:
        return []

    param_rows: List[Tuple[int, str]] = []
    blank_run = 0
    for r in range(marker_row + 1, end_row + 1):
        p = _param_match(ws_values.cell(r, best_col).value, known_params)
        if p:
            param_rows.append((r, p))
            blank_run = 0
        else:
            blank_run += 1
            if param_rows and blank_run >= 2:
                break

    if not param_rows:
        return []

    # Determine value columns. Ignore columns that contain status/current markers or repeat param names.
    candidate_cols = []
    for c in range(best_col + 1, end_col + 1):
        values = [ws_values.cell(r, c).value for r, _ in param_rows]
        nonblank = [v for v in values if not is_blank(v)]
        if not nonblank:
            continue
        if all(is_status(v) for v in nonblank):
            continue
        if any(CURRENT_RE.search(as_text(ws_values.cell(marker_row, c).value)) for _ in [0]):
            continue
        if all(_param_match(v, known_params) for v in nonblank if as_text(v)):
            continue
        candidate_cols.append(c)

    if not candidate_cols:
        return []

    # Stop candidate cols at the first obvious Current Values area.
    trimmed_cols = []
    for c in candidate_cols:
        marker_text = as_text(ws_values.cell(marker_row, c).value)
        if CURRENT_RE.search(marker_text):
            break
        trimmed_cols.append(c)
    candidate_cols = trimmed_cols or candidate_cols[:1]

    records: List[Dict[str, Any]] = []
    for c in candidate_cols:
        rec = {p: ws_values.cell(r, c).value for r, p in param_rows}
        if any(not is_blank(v) and not is_status(v) for v in rec.values()):
            if not record_contains_metadata_row(rec):
                records.append(rec)
    return records



def parse_embedded_parameter_target_values(
    ws_values,
    section: Section,
    next_cli_row: int,
) -> List[Dict[str, Any]]:
    """Parse a PARAMATERS table whose Target Value column is embedded in the table.

    Some CIQ sections do not have a separate "Target Values" block. Instead the
    PARAMATERS header contains a Target Value column and each parameter row has
    its desired value in that same column. Example:

        PARAMATERS | Type | Mandatory | CLI Name | ... | Target Value (Staging)
        mtc-conf-mtc-t3412-home | ... | [MTC_T3412_HOME] | ... | MME
        mtc-conf-mtc-t3412-roam | ... | [MTC_T3412_ROAM] | ... | MME

    This function converts the whole vertical parameter table into one target
    record so commands such as CHG-MMEMTC-CONF are generated as one combined
    TL1 command line.
    """
    max_row = min(ws_values.max_row, next_cli_row - 1 if next_cli_row else ws_values.max_row)

    header_row = None
    param_col = None
    cli_col = None
    target_col = None

    for r in range(section.cli_row, max_row + 1):
        row_texts = [norm_key(ws_values.cell(r, c).value) for c in range(1, scan_column_limit(ws_values) + 1)]
        if not any(t in {"parameters", "paramaters"} for t in row_texts):
            continue

        header_row = r
        for c in range(1, scan_column_limit(ws_values) + 1):
            raw = as_text(ws_values.cell(r, c).value)
            key = norm_key(raw)
            if key in {"parameters", "paramaters"}:
                param_col = c
            elif key == "cli name":
                cli_col = c
            elif TARGET_RE.search(raw):
                target_col = c
        break

    if header_row is None or cli_col is None or target_col is None:
        return []

    record: Dict[str, Any] = {}
    useful_rows = 0
    blank_run = 0

    for r in range(header_row + 1, max_row + 1):
        # Stop on a new CLI section or a true separate Target Values block.
        row_join = " | ".join(
            as_text(ws_values.cell(r, c).value)
            for c in range(1, min(ws_values.max_column, max(target_col, cli_col) + 2) + 1)
        )
        if CLI_RE.search(row_join):
            break

        cli_name, fixed_values = parse_cli_name_spec(ws_values.cell(r, cli_col).value)
        param = clean_param(ws_values.cell(r, param_col).value) if param_col else ""
        if is_placeholder_param(param):
            param = ""
        raw_target = ws_values.cell(r, target_col).value

        if not cli_name and not param and is_blank(raw_target):
            blank_run += 1
            if useful_rows and blank_run >= 2:
                break
            continue
        blank_run = 0

        if not cli_name:
            # If the explicit CLI Name cell is blank, try to infer it from the
            # parameter/YANG-style header.
            cli_name = infer_cli_from_param(param, section.cli_order) or ""

        if not cli_name:
            continue

        # Only consume parameters that belong to this command template.
        if section.cli_order and cli_name not in {x.upper() for x in section.cli_order}:
            inferred = infer_cli_from_param(param, section.cli_order)
            if inferred:
                cli_name = inferred
            else:
                continue

        # Ignore helper/status text rather than generating it into TL1.
        if is_metadata_or_helper_value(raw_target):
            continue

        # Use the parameter key when available so existing mapping still applies.
        # Also register direct CLI aliases defensively.
        key = param or cli_name
        section.param_to_cli.setdefault(key, cli_name)
        section.param_to_cli.setdefault(cli_name, cli_name)
        section.param_to_cli.setdefault(f"[{cli_name}]", cli_name)

        if fixed_values:
            if section.param_fixed_cli_values is None:
                section.param_fixed_cli_values = {}
            section.param_fixed_cli_values[key] = fixed_values

        record[key] = raw_target
        useful_rows += 1

    if not record:
        return []

    return [record]



def extract_act_lic_record(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract ACT-LIC information/signature target values.

    ACT-LIC uses labeled target rows instead of a normal horizontal/vertical
    target table, e.g.:
        Target Value (...) (information) | SN=...
        Target Value (...) (signature)   | MCw...
    """
    if section.command.upper() != "ACT-LIC":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    information = None
    signature = None

    for r in range(section.cli_row, min(next_cli_row - 1, ws.max_row) + 1):
        for c in range(1, scan_column_limit(ws) + 1):
            label = as_text(ws.cell(r, c).value)
            if not label:
                continue

            low = re.sub(r"\s+", " ", label).strip().lower()
            if not TARGET_RE.search(label):
                continue

            is_info = "information" in low
            is_sig = "signature" in low
            if not (is_info or is_sig):
                continue

            # Value is normally in the next non-empty cell on the same row.
            value = None
            for vc in range(c + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, vc).value
                if candidate is None or as_text(candidate) == "":
                    continue
                candidate_text = as_text(candidate)
                # Do not accidentally consume another label/status cell.
                if is_metadata_or_helper_value(candidate_text):
                    continue
                value = candidate_text
                break

            if value is None:
                continue

            if is_info and information is None:
                information = value
            elif is_sig and signature is None:
                signature = value

        if information is not None and signature is not None:
            break

    if information is None and signature is None:
        return []

    record: Dict[str, Any] = {}
    if information is not None:
        record["__ACT_LIC_INF__"] = information
    if signature is not None:
        record["__ACT_LIC_SIG__"] = signature
    return [record]


def quote_tl1_exact(value: Any) -> str:
    """Wrap a value in double quotes while otherwise preserving it exactly."""
    text = as_text(value)
    # Preserve embedded quotes safely for TL1 output.
    text = text.replace('"', '\\"')
    return f'"{text}"'


def records_to_act_lic_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render ACT-LIC as one exact-case INF/SIG command."""
    if section.command.upper() != "ACT-LIC" or not records:
        return []

    record = records[0]
    inf = record.get("__ACT_LIC_INF__")
    sig = record.get("__ACT_LIC_SIG__")

    parts: List[str] = []
    if inf is not None and as_text(inf) != "":
        parts.append(f"INF={quote_tl1_exact(inf)}")
    if sig is not None and as_text(sig) != "":
        parts.append(f"SIG={quote_tl1_exact(sig)}")

    if not parts:
        return []

    return [f"ACT-LIC:{','.join(parts)};"]



def extract_meas_kday_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Build MEAS-KDAY records from paired *-day / *-hour parameters.

    Some PM File Storage Period sections intentionally leave CLI Name blank and
    encode the command TYPE in the parameter name itself:

        raw-1min-day   -> TYPE=RAW-1MIN, DAYS=<target>
        raw-1min-hour  -> TYPE=RAW-1MIN, HOURS=<target>

    The same rule applies generically to any <type>-day / <type>-hour pair.
    """
    if section.command.upper() != "CHG-MEAS-KDAY":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    # Locate PARAMATERS header and its embedded Target Value column.
    header_row = None
    param_col = None
    target_col = None

    for r in range(section.cli_row, min(next_cli_row - 1, ws.max_row) + 1):
        row_texts = [norm_key(ws.cell(r, c).value) for c in range(1, scan_column_limit(ws) + 1)]
        if not any(t in {"parameters", "paramaters"} for t in row_texts):
            continue
        header_row = r
        for c in range(1, scan_column_limit(ws) + 1):
            raw = as_text(ws.cell(r, c).value)
            key = norm_key(raw)
            if key in {"parameters", "paramaters"}:
                param_col = c
            elif TARGET_RE.search(raw):
                target_col = c
        break

    if header_row is None or param_col is None or target_col is None:
        return []

    grouped: Dict[str, Dict[str, Any]] = {}
    order: List[str] = []

    for r in range(header_row + 1, min(next_cli_row - 1, ws.max_row) + 1):
        param = clean_param(ws.cell(r, param_col).value)
        if not param:
            continue

        m = re.fullmatch(r"(.+)-(day|hour)", param.strip(), re.I)
        if not m:
            # Once grouped rows have started, unrelated rows are documentation
            # or the next structure; ignore them rather than inventing output.
            continue

        type_name = m.group(1)
        kind = m.group(2).lower()
        raw_target = ws.cell(r, target_col).value

        if is_metadata_or_helper_value(raw_target):
            continue

        if type_name not in grouped:
            grouped[type_name] = {"TYPE": type_name}
            order.append(type_name)

        grouped[type_name]["DAYS" if kind == "day" else "HOURS"] = raw_target

    if not grouped:
        return []

    # Direct aliases allow generic record_to_cli_values/render_cli_line reuse.
    for cli in ("TYPE", "DAYS", "HOURS"):
        section.param_to_cli.setdefault(cli, cli)

    return [grouped[name] for name in order]




def _strip_single_value_braces(value: Any) -> str:
    """Remove outer { } from a single-value list representation."""
    text = as_text(value).strip()
    if text.startswith("{") and text.endswith("}"):
        return text[1:-1].strip()
    return text


def extract_csdb_local_remote_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Parse CRTE-CSDB:LOCALIP,LOCALPORT,[REMOTEIP],[REMOTEPORT],NAME.

    Supports CIQ layouts where:
      - a `local` row stores IP/port under csdb-info-ip / csdb-info-port
      - a `remote` row stores IP/port in the following two target columns
        (often labeled 0 / 0 or left blank in the template)
      - NAME is populated on each row

    Each source row becomes one TL1 record:
      local  -> LOCALIP, LOCALPORT, NAME
      remote -> REMOTEIP, REMOTEPORT, NAME
    """
    if section.command.upper() != "CRTE-CSDB":
        return []

    order = [x.upper() for x in section.cli_order]
    expected = ["LOCALIP", "LOCALPORT", "REMOTEIP", "REMOTEPORT", "NAME"]
    if order != expected:
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    markers = find_target_markers(ws, section.cli_row, next_cli_row - 1)
    if not markers:
        return []

    records: List[Dict[str, Any]] = []

    for marker_row, _marker_col in markers:
        header_row = marker_row + 1

        # Identify columns by their known YANG headers.
        header_by_col: Dict[int, str] = {}
        for c in range(1, scan_column_limit(ws) + 1):
            header = clean_param(ws.cell(header_row, c).value).strip().lower()
            if header:
                header_by_col[c] = header

        type_col = next((c for c, h in header_by_col.items()
                         if h == "csdb-info-single-redun-type"), None)
        ip_col = next((c for c, h in header_by_col.items()
                       if h == "csdb-info-ip"), None)
        port_col = next((c for c, h in header_by_col.items()
                         if h == "csdb-info-port"), None)
        name_col = next((c for c, h in header_by_col.items()
                         if h == "csdb-info-name"), None)

        if not all([type_col, ip_col, port_col, name_col]):
            continue

        # The remote IP/port columns immediately follow the local port column
        # in this layout, before NAME.
        remote_ip_col = port_col + 1
        remote_port_col = port_col + 2

        for r in range(header_row + 1, min(next_cli_row - 1, ws.max_row) + 1):
            redun_type = as_text(ws.cell(r, type_col).value).strip().lower()
            if not redun_type:
                # Stop after data has begun and a blank row is reached.
                if records:
                    break
                continue

            name = as_text(ws.cell(r, name_col).value).strip()
            rec: Dict[str, Any] = {}

            if redun_type == "local":
                ip = _strip_single_value_braces(ws.cell(r, ip_col).value)
                port = as_text(ws.cell(r, port_col).value).strip()
                if ip:
                    rec["LOCALIP"] = ip
                if port:
                    rec["LOCALPORT"] = port
                if name:
                    rec["NAME"] = name

            elif redun_type == "remote":
                ip = _strip_single_value_braces(ws.cell(r, remote_ip_col).value)
                port = as_text(ws.cell(r, remote_port_col).value).strip()

                # Defensive fallback for templates that keep remote values in
                # the same IP/PORT columns as local.
                if not ip:
                    ip = _strip_single_value_braces(ws.cell(r, ip_col).value)
                if not port:
                    port = as_text(ws.cell(r, port_col).value).strip()

                if ip:
                    rec["REMOTEIP"] = ip
                if port:
                    rec["REMOTEPORT"] = port
                if name:
                    rec["NAME"] = name
            else:
                continue

            if rec:
                records.append(rec)

        if records:
            break

    for cli in expected:
        section.param_to_cli.setdefault(cli, cli)

    return records


def extract_csdb_type_uri_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Parse the newer CRTE-CSDB:TYPE,URI target layout.

    Source table columns:
        csdb-info-single-redun-type
        csdb-info-ip
        csdb-info-port
        csdb-info-name

    The CLI only exposes TYPE and URI, so URI is composed as <IP>:<PORT>.
    Braces around leaf-list IP values are removed.
    """
    if section.command.upper() != "CRTE-CSDB":
        return []
    if [x.upper() for x in section.cli_order] != ["TYPE", "URI"]:
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    marker_rows = find_target_markers(ws, section.cli_row, next_cli_row - 1)
    if not marker_rows:
        return []

    records: List[Dict[str, Any]] = []
    for marker_row, _marker_col in marker_rows:
        header_row = marker_row + 1
        headers = {}
        for c in range(1, scan_column_limit(ws) + 1):
            h = clean_param(ws.cell(header_row, c).value).lower()
            if h:
                headers[h] = c

        type_col = headers.get("csdb-info-single-redun-type")
        ip_col = headers.get("csdb-info-ip")
        port_col = headers.get("csdb-info-port")
        if not (type_col and ip_col and port_col):
            continue

        for r in range(header_row + 1, min(next_cli_row - 1, ws.max_row) + 1):
            t = as_text(ws.cell(r, type_col).value)
            ip = as_text(ws.cell(r, ip_col).value)
            port = as_text(ws.cell(r, port_col).value)
            if not t and not ip and not port:
                break
            if not (t and ip and port):
                continue

            # Leaf-list single value may be written as {10.1.2.3}.
            ip = ip.strip()
            if ip.startswith("{") and ip.endswith("}"):
                ip = ip[1:-1].strip()

            records.append({"TYPE": t, "URI": f"{ip}:{port}"})

        if records:
            break

    section.param_to_cli.setdefault("TYPE", "TYPE")
    section.param_to_cli.setdefault("URI", "URI")
    return records








def _row_has_excluded_marker(
    ws_values,
    row: int,
    start_col: int,
    end_col: int,
) -> bool:
    """Detect an explicit excluded/not-selected marker on a target row."""
    bad = {"🗙", "✗", "✕", "×", "x"}
    for c in range(start_col, end_col + 1):
        value = as_text(ws_values.cell(row, c).value).strip().lower()
        if value in bad:
            return True
    return False



def extract_epc_snssai_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract EPC-SNSSAI parent rows and repeated child DNN/SNSSAI mappings."""
    if section.command.upper() != "CRTE-EPC-SNSSAI":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    all_records: List[Dict[str, Any]] = []

    target_params = {
        "epc-snssai-id",
        "epc-snssai-name",
        "epc-snssai-snssai-id",
        "epc-snssai-dnn-entry/epc-snssai-dnn",
        "epc-snssai-dnn-entry/epc-snssai-dnn-snssai-id",
    }

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        header_row = None
        header_items: List[Tuple[str, int]] = []

        for r in range(marker_row + 1, min(end_row, marker_row + 15) + 1):
            found: List[Tuple[str, int]] = []
            for c in range(1, scan_column_limit(ws) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in target_params:
                    found.append((p, c))

            if len(found) >= 3:
                id_cols = [c for p, c in found if p == "epc-snssai-id"]
                current_start_col = id_cols[1] if len(id_cols) > 1 else scan_column_limit(ws) + 1
                header_items = [(p, c) for p, c in found if c < current_start_col]
                if len(header_items) >= 3:
                    header_row = r
                    break

        if header_row is None:
            continue

        # Preserve first target-side occurrence of each parameter.
        target_cols: Dict[str, int] = {}
        for p, c in header_items:
            target_cols.setdefault(p, c)

        current_parent: Dict[str, Any] = {}

        def is_arrow(v: Any) -> bool:
            t = as_text(v).strip()
            return t in {"↪", "⇢"} or normalize_row_action(v) is not None

        for r in range(header_row + 1, end_row + 1):
            row_values = {
                p: ws.cell(r, c).value
                for p, c in target_cols.items()
            }

            id_raw = row_values.get("epc-snssai-id")
            name_raw = row_values.get("epc-snssai-name")
            parent_snssai_raw = row_values.get("epc-snssai-snssai-id")
            dnn_raw = row_values.get("epc-snssai-dnn-entry/epc-snssai-dnn")
            child_snssai_raw = row_values.get("epc-snssai-dnn-entry/epc-snssai-dnn-snssai-id")

            if all(is_blank(v) for v in row_values.values()):
                continue

            continuation = is_arrow(id_raw) or is_arrow(name_raw) or is_arrow(parent_snssai_raw)

            if not continuation:
                current_parent = {}

                if clean_cell_value(id_raw) is not None:
                    current_parent["epc-snssai-id"] = id_raw
                if clean_cell_value(name_raw) is not None:
                    current_parent["epc-snssai-name"] = name_raw
                if clean_cell_value(parent_snssai_raw) is not None:
                    current_parent["epc-snssai-snssai-id"] = parent_snssai_raw

                if "epc-snssai-id" not in current_parent:
                    continue

                rec = dict(current_parent)

                # Child-list values must be added separately after the create.
                if clean_cell_value(dnn_raw) is not None:
                    rec["__EPC_CHILD_DNN__"] = dnn_raw
                if clean_cell_value(child_snssai_raw) is not None:
                    rec["__EPC_CHILD_SNSSAI__"] = child_snssai_raw

                all_records.append(rec)
                continue

            if not current_parent:
                continue

            child_dnn = clean_cell_value(dnn_raw)
            child_snssai = clean_cell_value(child_snssai_raw)
            if child_dnn is None and child_snssai is None:
                continue

            rec = dict(current_parent)
            if child_dnn is not None:
                rec["__EPC_CHILD_DNN__"] = dnn_raw
            if child_snssai is not None:
                rec["__EPC_CHILD_SNSSAI__"] = child_snssai_raw
            rec["__ROW_ACTION__"] = "CHG"
            all_records.append(rec)

    section.param_to_cli.setdefault("epc-snssai-id", "ID")
    section.param_to_cli.setdefault("epc-snssai-name", "NAME")
    section.param_to_cli.setdefault("epc-snssai-snssai-id", "SNSSAI")

    return all_records


def extract_snssai_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract SNSSAI parent fields and repeated DNN child mappings.

    Every Target Values block yields:
      - one base profile context (ID/SST/SD/EBIALLOC/NASCGST/PRI)
      - one record per valid DNN/PDUCONF/INH child mapping

    Comparison/status markers such as X/cross do not suppress Target Values;
    DNN/PDUCONF/INH rows are generated whenever the target value itself is valid.
    """
    if section.command.upper() != "CRTE-SNSSAI":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    aliases = {
        "snssai-id": "ID",
        "snssai-sst": "SST",
        "snssai-sd": "SD",
        "snssai-ebialloc": "EBIALLOC",
        "snssai-nascgst": "NASCGST",
        "snssai-priority": "PRI",
        "snssai-dft-snssai": "DFT",
        "snssai-dnn-info-entry/snssai-dnn": "DNN",
        "snssai-dnn-info-entry/snssai-pduconf": "PDUCONF",
        "snssai-dnn-info-entry/snssai-inh": "INH",
    }

    base_cli = {"ID", "SST", "SD", "EBIALLOC", "NASCGST", "PRI", "DFT"}
    all_records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Find parameter-name column in this actual target block.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 120) + 1):
                if clean_param(ws.cell(r, c).value) in aliases:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        base_record: Dict[str, Any] = {}
        child_records: List[Dict[str, Any]] = []
        current_child: Optional[Dict[str, Any]] = None
        started = False
        blank_run = 0

        for r in range(marker_row + 1, end_row + 1):
            param = clean_param(ws.cell(r, param_col).value)
            cli = aliases.get(param)

            if not cli:
                row_text = " | ".join(
                    as_text(ws.cell(r, c).value)
                    for c in range(1, scan_column_limit(ws) + 1)
                    if not is_blank(ws.cell(r, c).value)
                )
                if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
                    if started:
                        break
                    continue
                if not row_text.strip():
                    if started:
                        blank_run += 1
                        if blank_run >= 2:
                            break
                    continue
                blank_run = 0
                continue

            started = True
            blank_run = 0

            # Read first target value to the right, before duplicated current
            # parameter/value columns begin.
            raw_value = None
            duplicate_param_col = None
            for c in range(param_col + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, c).value
                candidate_param = clean_param(candidate)
                if candidate_param in aliases:
                    duplicate_param_col = c
                    break
                if CURRENT_RE.search(as_text(candidate)):
                    duplicate_param_col = c
                    break
                if raw_value is None and not is_blank(candidate) and not is_status(candidate):
                    raw_value = candidate

            if raw_value is None:
                continue

            if clean_cell_value(raw_value) is None:
                continue

            # Comparison/status markers such as ✓ or 🗙 describe the current
            # state comparison. They do NOT suppress a Target Value. In
            # particular, a DNN row with 🗙 still needs to be generated so the
            # target configuration can be applied.
            if cli in base_cli:
                base_record[param] = raw_value
                continue

            if cli == "DNN":
                if current_child:
                    child_records.append(current_child)

                current_child = {param: raw_value}
                continue

            if cli in {"PDUCONF", "INH"}:
                if current_child is None:
                    continue
                current_child[param] = raw_value

        if current_child:
            child_records.append(current_child)

        # Emit one combined record per valid DNN mapping.
        for child in child_records:
            rec = dict(base_record)
            rec.update(child)
            all_records.append(rec)

    for param, cli in aliases.items():
        section.param_to_cli.setdefault(param, cli)

    return all_records


def extract_amf_peer_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract CRTE-AMF tables while preserving GUAMI, slice and service contexts.

    AMF Target Values mix several child lists:
      - parent: AMF, NAME, INSTID, ADMSTS, LOCALITY
      - GUAMI: MCC, MNC, AMFREG, AMFSET, AMFPTR
      - slice: SST, SD
      - service: SVC, SCHEME, SVCIP, SVCPORT

    The generic peer extractor cannot flatten these safely because repeated
    child-list rows would shift values into the wrong CLI fields.
    """
    if section.command.upper() != "CRTE-AMF":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    aliases = {
        "amf-id": "AMF",
        "amf-name": "NAME",
        "amf-inst-id": "INSTID",
        "amf-admsts": "ADMSTS",
        "amf-locality": "LOCALITY",
        "amf-guami-entry/amf-mcc": "MCC",
        "amf-guami-entry/amf-mnc": "MNC",
        "amf-guami-entry/amf-region-id": "AMFREG",
        "amf-guami-entry/amf-set-id": "AMFSET",
        "amf-guami-entry/amf-pointer": "AMFPTR",
        "amf-sliceinfo/amf-sst": "SST",
        "amf-sliceinfo/amf-sd": "SD",
        "amf-svc-entry/amf-svc-name": "SVC",
        "amf-svc-entry/amf-svc-scheme": "SCHEME",
        "amf-svc-entry/amf-svc-ipaddr-entry/amf-svc-ipaddr-ip": "SVCIP",
        "amf-svc-entry/amf-svc-ipaddr-entry/amf-svc-ipaddr-port": "SVCPORT",
    }

    all_records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Locate the parameter-name column in this actual Target Values block.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 120) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in aliases:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        parent: Dict[str, Any] = {}
        guami: Dict[str, Any] = {}
        slice_info: Dict[str, Any] = {}
        current_svc = None
        current_scheme = None
        current_ip = None
        current_port = None
        block_records: List[Dict[str, Any]] = []

        started = False
        blank_run = 0

        def set_param_value(rec: Dict[str, Any], cli_name: str, value: Any) -> None:
            for p, cli in aliases.items():
                if cli == cli_name:
                    rec[p] = value
                    return

        def flush_service_record() -> None:
            nonlocal current_ip, current_port
            if current_svc is None or current_ip is None:
                return

            rec: Dict[str, Any] = {}
            rec.update(parent)
            rec.update(guami)
            rec.update(slice_info)

            set_param_value(rec, "SVC", current_svc)
            if current_scheme is not None:
                set_param_value(rec, "SCHEME", current_scheme)
            set_param_value(rec, "SVCIP", current_ip)
            if current_port is not None:
                set_param_value(rec, "SVCPORT", current_port)

            if block_records:
                rec["__ROW_ACTION__"] = "CHG"
            block_records.append(rec)

            current_ip = None
            current_port = None

        for r in range(marker_row + 1, end_row + 1):
            param = clean_param(ws.cell(r, param_col).value)
            cli = aliases.get(param)

            if not cli:
                row_text = " | ".join(
                    as_text(ws.cell(r, c).value)
                    for c in range(1, scan_column_limit(ws) + 1)
                    if not is_blank(ws.cell(r, c).value)
                )
                if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
                    if started:
                        break
                    continue
                if not row_text.strip():
                    if started:
                        blank_run += 1
                        if blank_run >= 2:
                            break
                    continue
                blank_run = 0
                continue

            started = True
            blank_run = 0

            # Read target value to the right, stopping before Current Values.
            raw_value = None
            for c in range(param_col + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, c).value
                if is_blank(candidate) or is_status(candidate):
                    continue
                if clean_param(candidate) in aliases:
                    break
                if CURRENT_RE.search(as_text(candidate)):
                    break
                raw_value = candidate
                break

            if raw_value is None or clean_cell_value(raw_value) is None:
                continue

            if cli in {"AMF", "NAME", "INSTID", "ADMSTS", "LOCALITY"}:
                set_param_value(parent, cli, raw_value)
                continue

            if cli in {"MCC", "MNC", "AMFREG", "AMFSET", "AMFPTR"}:
                set_param_value(guami, cli, raw_value)
                continue

            if cli in {"SST", "SD"}:
                set_param_value(slice_info, cli, raw_value)
                continue

            if cli == "SVC":
                flush_service_record()
                current_svc = raw_value
                current_scheme = None
                continue

            if cli == "SCHEME":
                current_scheme = raw_value
                continue

            if cli == "SVCIP":
                if current_ip is not None:
                    flush_service_record()
                current_ip = raw_value
                continue

            if cli == "SVCPORT":
                current_port = raw_value
                flush_service_record()
                continue

        flush_service_record()
        if block_records:
            all_records.extend(block_records)

    for param, cli in aliases.items():
        section.param_to_cli.setdefault(param, cli)

    return all_records


def extract_peer_service_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract CRTE-<peer> service/IP child rows sequentially.

    Handles UDM/AUSF/NRF/SMF/PCF-style tables where parent identity fields are
    followed by repeated SVC -> SCHEME -> SVCIP -> SVCPORT sequences.

    Each emitted record is complete enough for rendering:
      - first IP of first service includes parent + SVC + SCHEME + SVCIP + port
      - additional IPs retain SVC and omit repeated SCHEME later at render time
      - new service starts a new SVC/SCHEME context
    """
    cmd = section.command.upper()
    if not cmd.startswith("CRTE-") or cmd == "CRTE-NE-GRP":
        return []
    if not _service_fields_present(section):
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    # Build param -> CLI aliases only for fields actually used by this command.
    aliases: Dict[str, str] = {}
    for param, cli in section.param_to_cli.items():
        if not isinstance(param, str):
            continue
        cli_u = (cli or "").upper()
        if cli_u in set(section.cli_order):
            aliases[clean_param(param)] = cli_u

    service_cli = {"SVC", "SCHEME", "SVCIP", "SVCPORT", "SVCID"}
    object_name = cmd[5:]
    key_candidates = {object_name, "ID"}

    parent_cli = {
        cli.upper() for cli in section.cli_order
        if cli.upper() not in service_cli
    }

    all_records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Find the parameter-name column in this actual Target Values block.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 150) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in aliases:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        parent_record: Dict[str, Any] = {}
        current_svc = None
        current_scheme = None
        current_ip = None
        current_port = None

        block_records: List[Dict[str, Any]] = []
        started = False
        blank_run = 0

        def flush_ip_record() -> None:
            nonlocal current_ip, current_port
            if current_svc is None or current_ip is None:
                return

            rec: Dict[str, Any] = dict(parent_record)

            # Store using the original mapped parameter key when possible.
            for p, cli in aliases.items():
                if cli == "SVC":
                    rec[p] = current_svc
                    break
            if current_scheme is not None:
                for p, cli in aliases.items():
                    if cli == "SCHEME":
                        rec[p] = current_scheme
                        break
            for p, cli in aliases.items():
                if cli == "SVCIP":
                    rec[p] = current_ip
                    break
            if current_port is not None:
                for p, cli in aliases.items():
                    if cli == "SVCPORT":
                        rec[p] = current_port
                        break

            # Mark continuation records after the first complete child row.
            if block_records:
                rec["__ROW_ACTION__"] = "CHG"
            block_records.append(rec)
            current_ip = None
            current_port = None

        for r in range(marker_row + 1, end_row + 1):
            param = clean_param(ws.cell(r, param_col).value)
            cli = aliases.get(param)

            if not cli:
                row_text = " | ".join(
                    as_text(ws.cell(r, c).value)
                    for c in range(1, scan_column_limit(ws) + 1)
                    if not is_blank(ws.cell(r, c).value)
                )
                if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
                    if started:
                        break
                    continue
                if not row_text.strip():
                    if started:
                        blank_run += 1
                        if blank_run >= 2:
                            break
                    continue
                blank_run = 0
                continue

            started = True
            blank_run = 0

            # Read target value to the right, before duplicated Current Values.
            raw_value = None
            for c in range(param_col + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, c).value
                if is_blank(candidate) or is_status(candidate):
                    continue
                if clean_param(candidate) in aliases:
                    break
                if CURRENT_RE.search(as_text(candidate)):
                    break
                raw_value = candidate
                break

            if raw_value is None:
                continue

            if clean_cell_value(raw_value) is None:
                continue

            if cli in parent_cli:
                parent_record[param] = raw_value
                continue

            if cli == "SVC":
                # Starting a new service means any pending IP record belongs
                # to the previous service and should be flushed first.
                flush_ip_record()
                current_svc = raw_value
                current_scheme = None
                continue

            if cli == "SCHEME":
                current_scheme = raw_value
                continue

            if cli == "SVCIP":
                # A new IP starts another child entry. Flush a previous pending
                # IP if it already had a port.
                if current_ip is not None:
                    flush_ip_record()
                current_ip = raw_value
                continue

            if cli == "SVCPORT":
                current_port = raw_value
                flush_ip_record()
                continue

        flush_ip_record()

        if block_records:
            all_records.extend(block_records)

    return all_records


def extract_ni_pf_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract CRTE-NI-PF Target Values preserving arrow continuation rows.

    A numeric ID row starts a new profile. Rows whose ID is an arrow inherit
    the previous ID/NAME and add another INTF_TYPE/NI pair.
    """
    if section.command.upper() != "CRTE-NI-PF":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    # Skip PARAMATERS table's embedded Target Value column.
    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    aliases = {
        "ni-profile-id": "ID",
        "ni-profile-name": "NAME",
        "ni-profile-ni-entry/ni-profile-intf-type": "INTF_TYPE",
        "ni-profile-ni-entry/ni-profile-ni": "NI",
    }

    records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Find header row that contains all/most NI-PF parameter names.
        header_row = None
        col_map: Dict[str, int] = {}

        for r in range(marker_row + 1, min(end_row, marker_row + 6) + 1):
            matches: Dict[str, int] = {}
            for c in range(1, scan_column_limit(ws) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in aliases:
                    matches[p] = c
            if len(matches) >= 3:
                header_row = r
                col_map = matches
                break

        if header_row is None:
            continue

        current_id = None
        current_name = None
        blank_run = 0

        for r in range(header_row + 1, end_row + 1):
            raw_id = ws.cell(r, col_map["ni-profile-id"]).value if "ni-profile-id" in col_map else None
            raw_name = ws.cell(r, col_map["ni-profile-name"]).value if "ni-profile-name" in col_map else None
            raw_intf = ws.cell(r, col_map["ni-profile-ni-entry/ni-profile-intf-type"]).value if "ni-profile-ni-entry/ni-profile-intf-type" in col_map else None
            raw_ni = ws.cell(r, col_map["ni-profile-ni-entry/ni-profile-ni"]).value if "ni-profile-ni-entry/ni-profile-ni" in col_map else None

            if all(is_blank(v) for v in (raw_id, raw_name, raw_intf, raw_ni)):
                blank_run += 1
                if blank_run >= 2:
                    break
                continue
            blank_run = 0

            # Ignore duplicated Current Values area or metadata rows.
            if any(is_metadata_or_helper_value(v) for v in (raw_id, raw_name, raw_intf, raw_ni)):
                continue

            if raw_id is not None and not is_arrow_value(raw_id) and not is_blank(raw_id):
                current_id = raw_id
            if raw_name is not None and not is_arrow_value(raw_name) and not is_blank(raw_name):
                current_name = raw_name

            if current_id is None:
                continue

            rec: Dict[str, Any] = {"ni-profile-id": current_id}

            if current_name is not None:
                rec["ni-profile-name"] = current_name

            if raw_intf is not None and not is_arrow_value(raw_intf) and not is_blank(raw_intf):
                rec["ni-profile-ni-entry/ni-profile-intf-type"] = raw_intf

            if raw_ni is not None and not is_arrow_value(raw_ni) and not is_blank(raw_ni):
                rec["ni-profile-ni-entry/ni-profile-ni"] = raw_ni

            # Require at least one child mapping to avoid emitting header-like rows.
            if (
                "ni-profile-ni-entry/ni-profile-intf-type" in rec
                and "ni-profile-ni-entry/ni-profile-ni" in rec
            ):
                records.append(rec)

        if records:
            break

    for param, cli in aliases.items():
        section.param_to_cli.setdefault(param, cli)

    return records


def extract_ne_grp_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract NE-GRP parent groups and repeated child members correctly.

    Each Target Values block becomes one parent record plus one child record
    per NEID. Repeated PRI/CAPA values are attached to the current NEID.
    """
    if section.command.upper() != "CRTE-NE-GRP":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    aliases = {
        "ne-group-type": "NE",
        "ne-group-id": "GID",
        "ne-group-name": "GNAME",
        "ne-group-bkup-grp-id": "BKUPGID",
        "ne-group-sel-mth": "SELM",
        "ne-group-tmr": "TMR",
        "ne-group-maxchg": "MAXCHG",
        "ne-group-af": "AF",
        "ne-group-elem-entry/ne-group-ne-id": "NEID",
        "ne-group-elem-entry/ne-group-pri": "PRI",
        "ne-group-elem-entry/ne-group-capa": "CAPA",
    }
    parent_cli = {"NE", "GID", "GNAME", "BKUPGID", "SELM", "TMR", "MAXCHG", "AF"}
    child_cli = {"NEID", "PRI", "CAPA"}

    all_records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Locate the parameter-name column in this actual Target Values block.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 100) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in aliases:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        parent_record: Dict[str, Any] = {}
        child_records: List[Dict[str, Any]] = []
        current_child: Optional[Dict[str, Any]] = None
        started = False
        blank_run = 0

        for r in range(marker_row + 1, end_row + 1):
            param = clean_param(ws.cell(r, param_col).value)

            if param not in aliases:
                # Unsupported NE-GRP parameters (for example ne-group-minu,
                # ne-group-grpadm, ne-group-use) can legitimately appear
                # between supported parent fields and child NEID rows.
                # Do not terminate on those rows. Terminate only on genuinely
                # blank separators or a new Target Values / section marker.
                row_values = [
                    ws.cell(r, c).value
                    for c in range(1, scan_column_limit(ws) + 1)
                ]
                row_text = " | ".join(as_text(v) for v in row_values if not is_blank(v))

                if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
                    if started:
                        break
                    continue

                if not row_text.strip():
                    if started:
                        blank_run += 1
                        if blank_run >= 2:
                            break
                    continue

                # Nonblank unsupported row inside this target block: ignore it
                # without advancing the blank-row termination counter.
                blank_run = 0
                continue

            started = True
            blank_run = 0
            cli = aliases[param]

            # Read target value to the right, before duplicated Current Values.
            raw_value = None
            for c in range(param_col + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, c).value
                if is_blank(candidate) or is_status(candidate):
                    continue

                candidate_param = clean_param(candidate)
                if candidate_param in aliases:
                    break
                if CURRENT_RE.search(as_text(candidate)):
                    break

                raw_value = candidate
                break

            if raw_value is None:
                continue

            # [ENC], SKIP, etc. must not be emitted.
            cleaned_value = clean_cell_value(raw_value)
            if cleaned_value is None:
                continue

            if cli in parent_cli:
                parent_record[param] = raw_value
                continue

            if cli == "NEID":
                if current_child:
                    child_records.append(current_child)
                current_child = {param: raw_value}
                continue

            if cli in child_cli:
                if current_child is None:
                    current_child = {}
                current_child[param] = raw_value

        if current_child:
            child_records.append(current_child)

        if parent_record:
            all_records.append(parent_record)
            all_records.extend(child_records)

    for param, cli in aliases.items():
        section.param_to_cli.setdefault(param, cli)

    return all_records


def extract_gtpc_cause_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract GTPC-CAUSE Target Values sequentially without collapsing repeats."""
    if section.command.upper() != "CRTE-GTPC-CAUSE":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    # IMPORTANT: skip the PARAMATERS header's "Target Value" column and start
    # searching only after the parameter-definition table.
    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    if not markers:
        return []

    known_params = {
        "gtpc-cause-index",
        "gtpc-cause-name",
        "gtpc-interface-entry/gtpc-cause-interface",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-protocol",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-error-code",
    }

    records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Locate the vertical parameter-name column in the actual Target Values block.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 100) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in known_params:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        started = False
        blank_run = 0
        for r in range(marker_row + 1, end_row + 1):
            param = clean_param(ws.cell(r, param_col).value)
            if param not in known_params:
                if started:
                    blank_run += 1
                    if blank_run >= 2:
                        break
                continue

            started = True
            blank_run = 0

            # Target value is the first meaningful cell to the right of param,
            # before duplicated Current Values columns begin.
            value = None
            for c in range(param_col + 1, scan_column_limit(ws) + 1):
                candidate = ws.cell(r, c).value
                if is_blank(candidate) or is_status(candidate):
                    continue

                candidate_param = clean_param(candidate)
                if candidate_param in known_params:
                    # This is the duplicated Current Values parameter column.
                    break
                if CURRENT_RE.search(as_text(candidate)):
                    break

                value = candidate
                break

            if value is not None:
                records.append({param: value})

        if records:
            break

    for param, cli in {
        "gtpc-cause-index": "INDEX",
        "gtpc-cause-name": "NAME",
        "gtpc-interface-entry/gtpc-cause-interface": "INTF",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-protocol": "PROTC",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-error-code": "GTPCC",
    }.items():
        section.param_to_cli.setdefault(param, cli)

    return records





def parameter_table_entries(
    ws_values,
    section: Section,
    next_cli_row: int,
) -> List[Tuple[str, str, str]]:
    """Return (parameter, type, cli-name) in Parameters-table order."""
    max_row = min(ws_values.max_row, next_cli_row - 1)
    header_row = None
    param_col = None
    type_col = None
    cli_col = None

    for r in range(section.cli_row, max_row + 1):
        for c in range(1, scan_column_limit(ws_values) + 1):
            key = norm_key(ws_values.cell(r, c).value)
            if key in {"parameters", "paramaters"}:
                header_row = r
                param_col = c
                break
        if header_row is not None:
            break

    if header_row is None or param_col is None:
        return []

    for c in range(1, scan_column_limit(ws_values) + 1):
        key = norm_key(ws_values.cell(header_row, c).value)
        if key == "type":
            type_col = c
        elif key == "cli name":
            cli_col = c

    entries: List[Tuple[str, str, str]] = []
    for r in range(header_row + 1, max_row + 1):
        row_text = " | ".join(
            as_text(ws_values.cell(r, c).value)
            for c in range(1, scan_column_limit(ws_values) + 1)
        )
        if TARGET_RE.search(row_text) or CLI_RE.search(row_text):
            break

        p = clean_param(ws_values.cell(r, param_col).value)
        if not p or is_placeholder_param(p):
            continue

        typ = as_text(ws_values.cell(r, type_col).value) if type_col else ""
        cli = clean_cli_name(ws_values.cell(r, cli_col).value) if cli_col else ""
        entries.append((p, typ, cli))

    return entries


def target_value_from_vertical_row(
    ws_values,
    row: int,
    param_col: int,
    known_params: set[str],
) -> Any:
    """Read the desired target value before the duplicated Current Values area."""
    raw_value = None
    for c in range(param_col + 1, scan_column_limit(ws_values) + 1):
        candidate = ws_values.cell(row, c).value
        if is_blank(candidate) or is_status(candidate):
            continue
        if clean_param(candidate) in known_params:
            break
        if CURRENT_RE.search(as_text(candidate)):
            break
        raw_value = candidate
        break
    return raw_value


def leading_row_has_arrow(ws_values, row: int, param_col: int) -> bool:
    """Detect ↪ / ⇢ group markers in cells to the left of the parameter."""
    for c in range(1, max(1, param_col)):
        if is_arrow_value(ws_values.cell(row, c).value):
            return True
        s = as_text(ws_values.cell(row, c).value)
        if "↪" in s or "⇢" in s:
            return True
    return False


def extract_pdu_conf_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract PDU-CONF parent records and DSCP children robustly.

    Recovery behavior:
      * each real Target Values block is processed independently;
      * parent rows are aligned to Parameters-table order;
      * if one source parameter label is missing/shifted, following values stay
        aligned positionally rather than shifting into the wrong CLI fields;
      * DSCP child groups are reconstructed from ↪/⇢ row grouping, so even
        damaged child parameter labels do not lose FQI/UL/DL values.
    """
    if section.command.upper() != "CRTE-PDU-CONF":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    section.param_to_cli["pdu-conf-dscp/pdu-conf-dscp-fqi"] = "DSCPFQI"
    section.param_to_cli["pdu-conf-dscp/pdu-conf-dscp-ul"] = "DSCPUL"
    section.param_to_cli["pdu-conf-dscp/pdu-conf-dscp-dl"] = "DSCPDL"

    entries = parameter_table_entries(ws, section, next_cli_row)
    parent_order = [
        p for p, typ, _cli in entries
        if "child list" not in norm_key(typ)
        and not p.startswith("pdu-conf-dscp/")
    ]
    known_params = set(section.param_to_cli.keys()) | {p for p, _, _ in entries}

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    results: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Locate the parameter-name column.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 180) + 1):
                if clean_param(ws.cell(r, c).value) in known_params:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None or best_count < 2:
            continue

        # Find first populated parent row, normally pdu-conf-id.
        data_rows: List[int] = []
        for r in range(marker_row + 1, end_row + 1):
            raw = target_value_from_vertical_row(ws, r, param_col, known_params)
            # Keep explicit placeholder rows such as [ENC] in the sequence so
            # later values do not shift into earlier parameters. They advance
            # positional alignment but are still omitted from TL1 output.
            if raw is None:
                continue
            data_rows.append(r)

        if not data_rows:
            continue

        parent: Dict[str, Any] = {}
        children: List[Dict[str, Any]] = []

        parent_idx = 0
        child_started = False
        current_child: Optional[Dict[str, Any]] = None
        child_pos = 0

        for r in data_rows:
            observed = clean_param(ws.cell(r, param_col).value)
            raw_value = target_value_from_vertical_row(ws, r, param_col, known_params)
            if raw_value is None:
                continue
            cleaned_value = clean_cell_value(raw_value)

            has_arrow = leading_row_has_arrow(ws, r, param_col)

            # DSCP groups are explicitly delimited by arrow rows. Once child
            # processing starts, map values by group position rather than by
            # potentially damaged labels.
            if has_arrow:
                child_started = True
                marker_text = " ".join(
                    as_text(ws.cell(r, c).value)
                    for c in range(1, max(2, param_col))
                )
                if "↪" in marker_text:
                    if current_child:
                        children.append(current_child)
                    current_child = {}
                    child_pos = 0

                if current_child is None:
                    current_child = {}

                if cleaned_value is not None:
                    if child_pos == 0:
                        current_child["pdu-conf-dscp/pdu-conf-dscp-fqi"] = raw_value
                    elif child_pos == 1:
                        current_child["pdu-conf-dscp/pdu-conf-dscp-ul"] = raw_value
                    elif child_pos == 2:
                        current_child["pdu-conf-dscp/pdu-conf-dscp-dl"] = raw_value
                child_pos += 1
                continue

            if child_started:
                continue

            # Ignore a pre-child [ENC] child placeholder row.
            if observed.startswith("pdu-conf-dscp/"):
                continue

            # Parent positional alignment. Normal blocks match exactly. If a
            # label is missing (as in the supplied ID=2000 block), this keeps
            # all subsequent values attached to the intended parameter order.
            if parent_idx >= len(parent_order):
                continue

            expected_param = parent_order[parent_idx]
            if cleaned_value is not None:
                parent[expected_param] = raw_value
            parent_idx += 1

        if current_child:
            children.append(current_child)

        if not parent:
            continue

        # Require a usable ID; empty placeholder blocks are ignored.
        id_param = next(
            (p for p, cli in section.param_to_cli.items()
             if cli.upper() == "ID" and p in parent),
            None,
        )
        if not id_param or clean_cell_value(parent.get(id_param)) in {None, ""}:
            continue

        results.append(parent)

        for child in children:
            rec: Dict[str, Any] = {
                id_param: parent[id_param],
                **child,
                "__ROW_ACTION__": "CHG",
            }
            results.append(rec)

    return results



def section_has_unmapped_pdu_dscp_children(
    wb_values,
    section: Section,
    sections: List[Section],
) -> bool:
    """PDU-CONF DSCP child mappings are supported."""
    return False



def extract_pcf_pf_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract PCF-PF parent and repeated child lists with positional recovery."""
    if section.command.upper() != "CRTE-PCF-PF":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    entries = parameter_table_entries(ws, section, next_cli_row)
    parent_order = [p for p, typ, _ in entries if "child list" not in norm_key(typ)]
    known_params = set(section.param_to_cli.keys()) | {p for p, _, _ in entries}

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    results: List[Dict[str, Any]] = []

    fail_params = [
        "pcf-pf-fail-actionlist/pcf-pf-fail-oper",
        "pcf-pf-fail-actionlist/pcf-pf-fail-oos",
        "pcf-pf-fail-actionlist/pcf-pf-fail-tout",
    ]
    event_params = [
        "pcf-pf-event-reportlist/pcf-pf-evrpt",
        "pcf-pf-event-reportlist/pcf-pf-evmode",
    ]

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 80) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in known_params or (p == 0):
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None:
            continue

        rows: List[Tuple[int, Any, Any]] = []
        for r in range(marker_row + 1, end_row + 1):
            raw = target_value_from_vertical_row(ws, r, param_col, known_params)
            if raw is None:
                continue
            rows.append((r, ws.cell(r, param_col).value, raw))

        if not rows:
            continue

        parent: Dict[str, Any] = {}
        fail_children: List[Dict[str, Any]] = []
        event_children: List[Dict[str, Any]] = []

        parent_idx = 0
        current_child: Optional[Dict[str, Any]] = None
        current_kind: Optional[str] = None
        child_pos = 0

        for r, observed_raw, raw_value in rows:
            observed = clean_param(observed_raw)
            lead_text = " ".join(
                as_text(ws.cell(r, c).value)
                for c in range(1, max(2, param_col))
            )
            has_arrow = "↪" in lead_text or "⇢" in lead_text or leading_row_has_arrow(ws, r, param_col)

            # Child continuation rows do not always repeat an arrow marker.
            # Once a child group has started, recognized child parameter rows
            # remain in that group until a new ↪ group begins.
            observed_is_fail_child = observed in fail_params
            observed_is_event_child = observed in event_params
            child_continuation = (
                current_child is not None
                and (
                    (current_kind == "fail" and observed_is_fail_child)
                    or (current_kind == "event" and observed_is_event_child)
                )
            )

            if has_arrow or child_continuation:
                if "↪" in lead_text:
                    if current_child:
                        if current_kind == "fail":
                            fail_children.append(current_child)
                        elif current_kind == "event":
                            event_children.append(current_child)
                    current_child = {}
                    child_pos = 0
                    current_kind = "event" if "(B)" in lead_text.upper() else "fail"

                if current_child is None:
                    current_child = {}

                params = event_params if current_kind == "event" else fail_params

                # Prefer actual child parameter identity when available. Fall
                # back to group position for damaged labels.
                if current_kind == "fail" and observed_is_fail_child:
                    target_param = observed
                elif current_kind == "event" and observed_is_event_child:
                    target_param = observed
                else:
                    target_param = params[child_pos] if child_pos < len(params) else None

                if target_param and clean_cell_value(raw_value) is not None:
                    current_child[target_param] = raw_value

                if target_param in params:
                    child_pos = max(child_pos + 1, params.index(target_param) + 1)
                else:
                    child_pos += 1
                continue

            # Parent rows are aligned positionally. This repairs the supplied
            # PFID=11 block where target-side parameter labels are shifted.
            if parent_idx < len(parent_order):
                if clean_cell_value(raw_value) is not None:
                    parent[parent_order[parent_idx]] = raw_value
                parent_idx += 1

        if current_child:
            if current_kind == "fail":
                fail_children.append(current_child)
            elif current_kind == "event":
                event_children.append(current_child)

        id_param = next(
            (p for p, cli in section.param_to_cli.items()
             if cli.upper() == "PFID" and p in parent),
            None,
        )
        if not id_param or clean_cell_value(parent.get(id_param)) in {None, ""}:
            continue

        # Parent record carries first valid child from each list when possible.
        parent_rec = dict(parent)
        if fail_children:
            parent_rec.update(fail_children[0])
        if event_children:
            parent_rec.update(event_children[0])
        results.append(parent_rec)

        # Additional child entries are preserved as explicit continuation rows.
        for child in fail_children[1:]:
            results.append({
                id_param: parent[id_param],
                **child,
                "__ROW_ACTION__": "CHG",
                "__PCF_CHILD_KIND__": "fail",
            })
        for child in event_children[1:]:
            results.append({
                id_param: parent[id_param],
                **child,
                "__ROW_ACTION__": "CHG",
                "__PCF_CHILD_KIND__": "event",
            })

    return results




def extract_onc_dict_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract ONC-DICT from its Parameters table target values.

    ONC-DICT commonly stores desired values directly in the Parameters table.
    The remarks/target helper cells may contain text such as "Refer to column B
    below.", which must never be rendered as ATTR/INIT/UPDATE/TERM values.
    """
    if section.command.upper() != "CHG-ONC-DICT":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    recs = parse_embedded_parameter_target_values(ws, section, next_cli_row)
    cleaned_recs: List[Dict[str, Any]] = []

    for rec in recs:
        out: Dict[str, Any] = {}
        for p, v in rec.items():
            if str(p).startswith("__"):
                out[p] = v
                continue
            if is_reference_helper_value(v):
                continue
            out[p] = v
        if out:
            cleaned_recs.append(out)

    return cleaned_recs




def extract_aprsc_lmt_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract application resource limit rows.

    Source table commonly contains:
        limit-id | resource-type | critical-alarm | critical-clear |
        major-alarm | major-clear | minor-alarm | minor-clear

    CLI:
        CRTE-APRSC-LMT:RSC,CRIA,CRIC,MAJA,MAJC,MINA,MINC,[DESC];

    `resource-type` is descriptive/source context and is intentionally not
    emitted because it has no CLI name in this command.
    """
    if section.command.upper() != "CRTE-APRSC-LMT":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    # Protect explicit mappings even if the Parameters table contains blanks.
    aliases = {
        "limit-id": "RSC",
        "critical-alarm": "CRIA",
        "critical-clear": "CRIC",
        "major-alarm": "MAJA",
        "major-clear": "MAJC",
        "minor-alarm": "MINA",
        "minor-clear": "MINC",
    }
    for p, cli in aliases.items():
        section.param_to_cli[p] = cli

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Find horizontal header row containing the relevant source parameters.
        header_row = None
        header_cols: Dict[str, int] = {}
        expected_headers = {
            "limit-id",
            "resource-type",
            "critical-alarm",
            "critical-clear",
            "major-alarm",
            "major-clear",
            "minor-alarm",
            "minor-clear",
        }

        for r in range(marker_row + 1, min(end_row, marker_row + 5) + 1):
            found: Dict[str, int] = {}
            for c in range(1, scan_column_limit(ws) + 1):
                p = clean_param(ws.cell(r, c).value).lower()
                if p in expected_headers and p not in found:
                    found[p] = c
            if "limit-id" in found and "critical-alarm" in found and "minor-clear" in found:
                header_row = r
                header_cols = found
                break

        if header_row is None:
            continue

        blank_run = 0
        for r in range(header_row + 1, end_row + 1):
            id_raw = ws.cell(r, header_cols["limit-id"]).value
            if is_blank(id_raw):
                blank_run += 1
                if blank_run >= 2:
                    break
                continue
            blank_run = 0

            id_clean = clean_cell_value(id_raw)
            if id_clean in {None, ""} or is_status(id_raw):
                continue
            if not re.fullmatch(r"-?\d+", as_text(id_clean)):
                continue

            rec: Dict[str, Any] = {"limit-id": id_raw}
            for param in (
                "critical-alarm",
                "critical-clear",
                "major-alarm",
                "major-clear",
                "minor-alarm",
                "minor-clear",
            ):
                col = header_cols.get(param)
                if col is None:
                    continue
                raw = ws.cell(r, col).value
                if raw is None:
                    continue
                if clean_cell_value(raw) is None:
                    continue
                rec[param] = raw

            # Optional DESC can be present in some variants. If available,
            # preserve it and let existing DESC case rules handle formatting.
            desc_col = header_cols.get("desc")
            if desc_col:
                raw = ws.cell(r, desc_col).value
                if raw is not None and clean_cell_value(raw) is not None:
                    rec["DESC"] = raw
                    section.param_to_cli["DESC"] = "DESC"

            records.append(rec)

    return records


def extract_supp_lmt_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract CHG-SUPP-LMT records from fault suppression target tables.

    Expected source shape:
        fault-code | inhibit | supp-limit

    CLI mapping:
        fault-code -> CODE
        supp-limit -> LMT

    The `inhibit` column is intentionally ignored because it is not part of
    CHG-SUPP-LMT:[CODE],LMT.
    """
    if section.command.upper() != "CHG-SUPP-LMT":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    section.param_to_cli["fault-code"] = "CODE"
    section.param_to_cli["supp-limit"] = "LMT"

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        header_row = None
        code_col = None
        lmt_col = None

        target_start_col, target_end_col = target_column_bounds(
            ws, marker_row, marker_col
        )

        for r in range(marker_row + 1, min(end_row, marker_row + 5) + 1):
            row_code_col = None
            row_lmt_col = None
            for c in range(target_start_col, target_end_col + 1):
                p = clean_param(ws.cell(r, c).value).lower()
                if p == "fault-code" and row_code_col is None:
                    row_code_col = c
                elif p == "supp-limit" and row_lmt_col is None:
                    row_lmt_col = c

            if row_code_col is not None and row_lmt_col is not None:
                header_row = r
                code_col = row_code_col
                lmt_col = row_lmt_col
                break

        if header_row is None:
            continue

        blank_run = 0
        for r in range(header_row + 1, end_row + 1):
            code_raw = ws.cell(r, code_col).value
            lmt_raw = ws.cell(r, lmt_col).value

            if is_blank(code_raw) and is_blank(lmt_raw):
                blank_run += 1
                if blank_run >= 2:
                    break
                continue
            blank_run = 0

            code_clean = clean_cell_value(code_raw)
            lmt_clean = clean_cell_value(lmt_raw)

            if code_clean in {None, ""} or lmt_clean in {None, ""}:
                continue
            if is_reference_helper_value(code_raw) or is_reference_helper_value(lmt_raw):
                continue

            records.append({
                "fault-code": code_raw,
                "supp-limit": lmt_raw,
            })

    return records


def extract_cha_rule_records(
    wb_values,
    section: Section,
    sections: List[Section],
) -> List[Dict[str, Any]]:
    """Extract one CHA-RULE record per real Target Values block.

    This parser is intentionally positional. If a target block has shifted or
    damaged parameter labels, values are aligned to the Parameters-table order
    rather than being assigned to the wrong CLI names.

    It also relies on real Target Values block boundaries, so a corrupted data
    row containing the literal text "Target Values (B1C3)" does not become a
    new configuration object.
    """
    if section.command.upper() != "CRTE-CHA-RULE":
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections
         if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
        or [ws.max_row + 1]
    )

    entries = parameter_table_entries(ws, section, next_cli_row)
    param_order = [p for p, _typ, _cli in entries]
    known_params = set(section.param_to_cli.keys()) | set(param_order)

    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    records: List[Dict[str, Any]] = []

    for marker_row, marker_col in markers:
        end_row = _region_end_row(ws, marker_row, next_cli_row)

        # Find the dominant target-side parameter-name column.
        param_col = None
        best_count = 0
        for c in range(1, scan_column_limit(ws) + 1):
            count = 0
            for r in range(marker_row + 1, min(end_row, marker_row + 40) + 1):
                p = clean_param(ws.cell(r, c).value)
                if p in known_params:
                    count += 1
            if count > best_count:
                best_count = count
                param_col = c

        if param_col is None:
            continue

        # Collect all rows containing a target value. Preserve [ENC] rows in
        # the sequence so following values do not shift to earlier parameters.
        target_rows: List[Tuple[int, Any]] = []
        for r in range(marker_row + 1, end_row + 1):
            raw = target_value_from_vertical_row(ws, r, param_col, known_params)
            if raw is None:
                continue
            target_rows.append((r, raw))

        if not target_rows:
            continue

        rec: Dict[str, Any] = {}
        pos = 0

        for r, raw_value in target_rows:
            if pos >= len(param_order):
                break

            expected_param = param_order[pos]
            cleaned = clean_cell_value(raw_value)

            # Store valid values, but always advance the positional slot.
            if cleaned is not None:
                rec[expected_param] = raw_value
            pos += 1

        # Require the key CHAID. Empty placeholder blocks are ignored.
        chaid_param = next(
            (p for p, cli in section.param_to_cli.items()
             if cli.upper() == "CHAID" and p in rec),
            None,
        )
        if not chaid_param:
            continue
        if clean_cell_value(rec.get(chaid_param)) in {None, ""}:
            continue

        records.append(rec)

    return records


def extract_target_records_for_section(wb_values, section: Section, sections: List[Section]) -> List[Dict[str, Any]]:
    if section.command.upper() == "ACT-LIC":
        return extract_act_lic_record(wb_values, section, sections)

    if section.command.upper() == "CRTE-PDU-CONF":
        pdu_records = extract_pdu_conf_records(wb_values, section, sections)
        if pdu_records:
            return pdu_records

    if section.command.upper() == "CRTE-APRSC-LMT":
        aprsc_records = extract_aprsc_lmt_records(wb_values, section, sections)
        if aprsc_records:
            return aprsc_records

    if section.command.upper() == "CHG-SUPP-LMT":
        supp_lmt_records = extract_supp_lmt_records(wb_values, section, sections)
        if supp_lmt_records:
            return supp_lmt_records

    if section.command.upper() == "CRTE-CHA-RULE":
        cha_rule_records = extract_cha_rule_records(wb_values, section, sections)
        if cha_rule_records:
            return cha_rule_records

    if section.command.upper() == "CHG-ONC-DICT":
        onc_dict_records = extract_onc_dict_records(wb_values, section, sections)
        if onc_dict_records:
            return onc_dict_records

    if section.command.upper() == "CRTE-PCF-PF":
        pcf_records = extract_pcf_pf_records(wb_values, section, sections)
        if pcf_records:
            return pcf_records

    if section.command.upper() == "CRTE-EPC-SNSSAI":
        epc_records = extract_epc_snssai_records(wb_values, section, sections)
        if epc_records:
            return epc_records

    if section.command.upper() == "CRTE-SNSSAI":
        snssai_records = extract_snssai_records(wb_values, section, sections)
        if snssai_records:
            return snssai_records

    if section.command.upper() == "CRTE-AMF":
        amf_records = extract_amf_peer_records(wb_values, section, sections)
        if amf_records:
            return amf_records

    if section.command.upper().startswith("CRTE-") and _service_fields_present(section):
        peer_records = extract_peer_service_records(wb_values, section, sections)
        if peer_records:
            return peer_records

    if section.command.upper() == "CRTE-NI-PF":
        ni_pf_records = extract_ni_pf_records(wb_values, section, sections)
        if ni_pf_records:
            return ni_pf_records

    if section.command.upper() == "CRTE-NE-GRP":
        ne_grp_records = extract_ne_grp_records(wb_values, section, sections)
        if ne_grp_records:
            return ne_grp_records

    if section.command.upper() == "CRTE-GTPC-CAUSE":
        gtpc_records = extract_gtpc_cause_records(wb_values, section, sections)
        if gtpc_records:
            return gtpc_records

    if section.command.upper() == "CRTE-CSDB":
        csdb_order = [x.upper() for x in section.cli_order]

        if csdb_order == ["LOCALIP", "LOCALPORT", "REMOTEIP", "REMOTEPORT", "NAME"]:
            csdb_records = extract_csdb_local_remote_records(wb_values, section, sections)
            if csdb_records:
                return csdb_records

        if csdb_order == ["TYPE", "URI"]:
            csdb_records = extract_csdb_type_uri_records(wb_values, section, sections)
            if csdb_records:
                return csdb_records

    if section.command.upper() == "CHG-MEAS-KDAY":
        meas_records = extract_meas_kday_records(wb_values, section, sections)
        if meas_records:
            return meas_records

    known_params = set(section.param_to_cli.keys())
    if not known_params:
        return []

    ws = wb_values[section.sheet_name]
    next_cli_row = min(
        [s.cli_row for s in sections if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row] or [ws.max_row + 1]
    )

    # 1) Try target markers in this section only.
    # Ignore the Target Value/Current Value columns that appear in the Parameters table header.
    param_table_end = find_parameter_table_end(ws, section, next_cli_row)
    markers = find_target_markers(ws, param_table_end + 1, next_cli_row - 1)
    records: List[Dict[str, Any]] = []
    for marker_row, marker_col in markers:
        block_records = parse_horizontal_targets(
            ws, marker_row, marker_col, known_params, next_cli_row
        )
        if not block_records:
            block_records = parse_vertical_targets(
                ws, marker_row, marker_col, known_params, next_cli_row
            )
        records.extend(block_records)

    # 2) If target block is only a pointer, follow "Refer to <sheet> sheet".
    if not records:
        referred = find_referred_sheet_name(ws, section.cli_row, next_cli_row - 1)
        referred_sheet_name = resolve_referred_sheet(wb_values, referred) if referred else None
        if referred_sheet_name:
            ref_ws = wb_values[referred_sheet_name]
            ref_markers = find_target_markers(ref_ws)
            for marker_row, marker_col in ref_markers:
                block_records = parse_horizontal_targets(
                    ref_ws, marker_row, marker_col, known_params
                )
                if not block_records:
                    block_records = parse_vertical_targets(
                        ref_ws, marker_row, marker_col, known_params
                    )
                records.extend(block_records)

    # 3) Fallback: some sections keep Target Value directly inside the
    # PARAMATERS table and have no separate Target Values block.
    if not records:
        records = parse_embedded_parameter_target_values(ws, section, next_cli_row)

    return records


# -----------------------------
# TL1 rendering
# -----------------------------

def is_metadata_or_helper_value(value: Any) -> bool:
    """Return True for CIQ header/status/helper cells that must never become TL1 values.

    This prevents output such as:
        CHG-AMFTA-INF:MNC="REVIEW STATUS",TAL=PENDING;
        CHG-AMFTA-INF:MCC="YANG PATH: ...";
    which can happen when a shifted Target Values marker causes the parser to
    treat section metadata rows as target data.
    """
    text = as_text(value).strip()
    if not text:
        return False

    low = re.sub(r"\s+", " ", text).strip().lower()

    exact_bad = {
        "review status",
        "config status",
        "description",
        "description:",
        "dependencies",
        "environment specific?",
        "parameters",
        "paramaters",
        "cli name",
        "remarks",
        "target value",
        "target values",
        "current value",
        "current values",
        "type",
        "mandatory",
        "details hidden",
        "resource details",
        "call timer details",
        "license details",
        "value range",
        "recommended",
    }
    if low in exact_bad:
        return True

    bad_prefixes = (
        "yang path:",
        "cli command:",
        "internal sanity check:",
        "target value ",
        "target values ",
        "current value ",
        "current values ",
        "resource details",
        "call timer details",
        "license details",
        "details hidden",
    )
    if low.startswith(bad_prefixes):
        return True

    # Header-like rows often contain these fragments and should be ignored when
    # they appear as target values.
    bad_fragments = (
        "review status",
        "config status",
        "yang path:",
        "cli command:",
    )
    return any(fragment in low for fragment in bad_fragments)


def record_contains_metadata_row(record: Dict[str, Any]) -> bool:
    """Detect and skip whole records accidentally parsed from CIQ metadata rows."""
    return any(is_metadata_or_helper_value(v) for v in record.values())


def record_to_cli_values(section: Section, record: Dict[str, Any], *, keep_arrows: bool = False) -> Dict[str, str]:
    """Build CLI-name -> formatted value from one target record."""
    if record_contains_metadata_row(record):
        return {}

    cli_values: Dict[str, str] = {}
    for param, raw_value in record.items():
        if param == "__ROW_ACTION__":
            continue
        cli_name = section.param_to_cli.get(param)

        # For MMECALL-PROF, the CIQ template sometimes has swapped/misaligned
        # CLI Name cells in the PARAMATERS table; the Target Values headers are
        # the reliable source. Force known call-prof-* headers to their correct
        # TL1 names before formatting. This prevents:
        #   call-prof-dcn_mmeg -> NEPSIDX
        #   call-prof-nepsidx  -> DCN_MMEG
        # and restores DCN_MMEG=NONE, NEPSIDX=1.
        if section.command.upper() == "CHG-MMECALL-PROF":
            inferred_cli = infer_cli_from_param(param, section.cli_order)
            if inferred_cli:
                cli_name = inferred_cli

        if not cli_name:
            continue
        if is_arrow_value(raw_value) and not keep_arrows:
            continue
        cleaned = clean_cell_value(raw_value)
        if cleaned is None:
            continue
        formatted = format_tl1_value(cli_name, cleaned, section.command)
        if cli_name.upper() == "NE" and section.command.upper() in {"CRTE-NE-GRP", "CHG-NE-GRP"}:
            formatted = normalize_ne_type(formatted)
        cli_values[cli_name] = formatted

    fixed_values = extract_fixed_cli_values(section.raw_cli)
    for fixed_name, fixed_value in fixed_values.items():
        cli_values[fixed_name] = fixed_value

    return cli_values


def render_cli_line(
    command: str,
    cli_order: Sequence[str],
    values: Dict[str, str],
    *,
    required_order: Optional[Sequence[str]] = None,
    keep_empty_fields: Optional[set[str]] = None,
) -> Optional[str]:
    """Render one TL1 line using the command/template order."""
    order = list(required_order or cli_order or values.keys())
    keep_empty_fields = {x.upper() for x in (keep_empty_fields or set())}
    parts: List[str] = []
    seen = set()
    for cli_name in order:
        cli_name = cli_name.upper()
        if cli_name in seen:
            continue
        seen.add(cli_name)
        if cli_name not in values:
            continue
        value = values[cli_name]

        # NE-GRP convention: do not emit BKUPGID when the configured value is 0.
        # Example:
        #   CRTE-NE-GRP:NE=UDM,GID=1,GNAME=LAB-UDM-GRP1,BKUPGID=0,SELM=CAPA,...
        # becomes:
        #   CRTE-NE-GRP:NE=UDM,GID=1,GNAME=LAB-UDM-GRP1,SELM=CAPA,...
        if command.upper() == "CRTE-NE-GRP" and cli_name == "BKUPGID" and str(value).strip() == "0":
            continue

        if value == "" and cli_name not in keep_empty_fields:
            continue
        parts.append(f"{cli_name}={value}")
    if not parts:
        return None
    return f"{command}:{','.join(parts)};"










def records_to_ofc_evt_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render only explicitly enabled/configured OFC event rows.

    Default-off rows are omitted. Fixed ITEM/EVT values from the CLI Name cell
    are retained. CAT is omitted when its source is [ENC]/unavailable.
    """
    if section.command.upper() != "CHG-OFC-EVT":
        return []

    lines: List[str] = []
    order = ["ITEM", "EVT", "ENABLE", "CAT"]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals:
            continue

        enable = vals.get("ENABLE")
        if not enable:
            continue

        # Skip rows whose target is only the default OFF setting.
        raw_non_control = [
            v for k, v in rec.items()
            if not str(k).startswith("__")
        ]
        if raw_non_control and all(
            is_default_wrapped_value(v) or clean_cell_value(v) is None
            for v in raw_non_control
        ):
            continue

        # The intended event outputs are enabled events only.
        if enable.upper() != "ON":
            continue

        line = render_cli_line(
            "CHG-OFC-EVT",
            order,
            vals,
            required_order=order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_onc_rcode_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render ONC result-code child entries with mandatory MODE=ADD.

    The source CLI template is:
        CHG-ONC-RCODE:ID,MODE,TYPE,[START],[END],[RACT];

    Target tables do not normally contain a MODE column, so MODE=ADD must be
    supplied by the generator for every result-code child entry.
    """
    if section.command.upper() != "CHG-ONC-RCODE":
        return []

    lines: List[str] = []
    order = ["ID", "MODE", "TYPE", "START", "END", "RACT"]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals:
            continue
        if "ID" not in vals or "TYPE" not in vals:
            continue

        out_vals: Dict[str, str] = {
            "ID": vals["ID"],
            "MODE": "ADD",
            "TYPE": vals["TYPE"],
        }
        for key in ("START", "END", "RACT"):
            if key in vals:
                out_vals[key] = vals[key]

        line = render_cli_line(
            "CHG-ONC-RCODE",
            order,
            out_vals,
            required_order=order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_rbase_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render one CRTE-RBASE per RBID and add repeated RULE children.

    Horizontal Target Values rows commonly use:
        first row: RBID, NAME, [PREC], [PDRSET], RULE
        later rows: ↪, ⇢, ⇢, ⇢, RULE

    The first valid RULE stays on CRTE-RBASE. Additional RULE values are
    emitted as CHG-RBASE with MODE=ADD so a parent rule base is not recreated
    for every child rule.
    """
    if section.command.upper() != "CRTE-RBASE":
        return []

    lines: List[str] = []
    created: set[str] = set()

    # Parent create follows the actual CRTE template. PREC may exist in the
    # parameter table but is intentionally omitted if it is not in cli_order.
    create_order = list(section.cli_order)
    chg_order = ["RBID", "MODE", "RULE"]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals or "RBID" not in vals:
            continue

        rbid = vals["RBID"]

        if rbid not in created:
            line = render_cli_line(
                "CRTE-RBASE",
                create_order,
                vals,
                required_order=create_order,
            )
            if line:
                lines.append(line)
                created.add(rbid)
            continue

        # Continuation rows only add RULE. Inherited NAME/PDRSET/etc. must not
        # be repeated on CHG-RBASE.
        if "RULE" not in vals:
            continue

        chg_vals = {
            "RBID": rbid,
            "MODE": "ADD",
            "RULE": vals["RULE"],
        }
        line = render_cli_line(
            "CHG-RBASE",
            chg_order,
            chg_vals,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_pcf_pf_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render PCF-PF parent plus preserved child-list continuation rows."""
    if section.command.upper() != "CRTE-PCF-PF":
        return []

    lines: List[str] = []
    created: set[str] = set()

    # Child-list continuation ordering. MODE is deliberately before child keys.
    fail_order = ["PFID", "MODE", "PFOPR", "PFOOS", "PFTOUT"]
    event_order = ["PFID", "MODE", "EVTRPT", "EVMODE"]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals or "PFID" not in vals:
            continue

        pfid = vals["PFID"]
        kind = rec.get("__PCF_CHILD_KIND__")

        if kind:
            chg_vals = {"PFID": pfid, "MODE": "ADD"}
            if kind == "fail":
                for k in ("PFOPR", "PFOOS", "PFTOUT"):
                    if k in vals:
                        chg_vals[k] = vals[k]
                line = render_cli_line(
                    "CHG-PCF-PF", fail_order, chg_vals, required_order=fail_order
                )
            else:
                for k in ("EVTRPT", "EVMODE"):
                    if k in vals:
                        chg_vals[k] = vals[k]
                line = render_cli_line(
                    "CHG-PCF-PF", event_order, chg_vals, required_order=event_order
                )
            if line:
                lines.append(line)
            continue

        if pfid in created:
            continue

        line = render_cli_line(
            "CRTE-PCF-PF",
            section.cli_order,
            vals,
            required_order=section.cli_order,
        )
        if line:
            lines.append(line)
            created.add(pfid)

    return lines


def records_to_pdu_conf_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render CRTE-PDU-CONF plus repeated CHG-PDU-CONF DSCP children."""
    if section.command.upper() != "CRTE-PDU-CONF":
        return []

    lines: List[str] = []
    created_ids: set[str] = set()

    chg_order = [
        "ID", "MODE", "NAME", "PDUTYPE", "PREFTYPE", "NONIP_TYPE", "NEW_PDUTYPE",
        "SSCMODE", "SELM", "LADN", "LADNACT", "NASCGST", "PDNS4", "SDNS4", "PDNS6",
        "SDNS6", "PCSCFPF", "MTU", "S6B_GRP_ID", "TUNNEL_UDP_PORT", "FWD4", "FWD6",
        "FQI", "ARP", "ARPCAP", "ARPVUL", "UAMBR", "DAMBR", "IWEPC", "SUBS_USE",
        "REDSALW", "PCFPF", "HF", "HCCPF", "RF", "RCCPF", "VF", "VCCPF", "PPI_SUPPORT",
        "REACT", "QOS_MARK", "ADNN", "IGUEIP", "IGSTATIC", "VNGRP", "ALWAYSON", "DELIWK",
        "PCF_PRE_IWK", "INACT_SUPP", "INACT_ACT", "INACT_TMR", "ETH_INACT_TMR",
        "LCALL_TMR", "NIPF", "BUF_CNT", "INACT_ACT_PCF_F", "INACT_TMR_PCF_F",
        "LCALL_TMR_PCF_F", "DFLT_TMR_SUPP", "DEDI_TMR_SUPP", "ICMP_MTU", "RA_INTERVAL",
        "RA_COUNT", "RA_LIFETIME", "RA_IDLE_IGN", "BCM", "RVAL", "CRACT", "CRACT_MODE",
        "RSTOPF", "UPSSUPP", "INTEG", "MAXUL", "RELPREF", "CIPHER", "MAXDL", "ACCT",
        "ACCT_MODE", "RACCT_ID", "RACCT_URR", "MIRROR_ID", "ACCTINTRM", "ACCTINTV",
        "AUTH", "AUTH_MODE", "AUTHINFO", "AUTHNAME", "AUTH_PASSWD", "RAUTH_ID", "AUTHBYP",
        "LCL_QOS", "LQACT", "QMAP", "DSCPFQI", "DSCPUL", "DSCPDL", "DESC"
    ]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals or "ID" not in vals:
            continue

        pdu_id = vals["ID"]
        is_child = any(k in vals for k in ("DSCPFQI", "DSCPUL", "DSCPDL"))

        if not is_child:
            if pdu_id in created_ids:
                continue
            line = render_cli_line(
                "CRTE-PDU-CONF",
                section.cli_order,
                vals,
                required_order=section.cli_order,
            )
            if line:
                lines.append(line)
                created_ids.add(pdu_id)
            continue

        chg_vals: Dict[str, str] = {
            "ID": pdu_id,
            "MODE": "ADD",
        }
        for k in ("DSCPFQI", "DSCPUL", "DSCPDL"):
            if k in vals:
                chg_vals[k] = vals[k]

        line = render_cli_line(
            "CHG-PDU-CONF",
            chg_order,
            chg_vals,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_epc_snssai_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render EPC-SNSSAI parent creates and child additions separately."""
    if section.command.upper() != "CRTE-EPC-SNSSAI":
        return []

    lines: List[str] = []
    created_ids: set[str] = set()

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals or "ID" not in vals:
            continue

        epc_id = vals["ID"]

        child_dnn = None
        child_snssai = None

        child_dnn_raw = rec.get("__EPC_CHILD_DNN__")
        if child_dnn_raw is not None and clean_cell_value(child_dnn_raw) is not None:
            child_dnn = format_tl1_value(
                "DNN",
                as_text(child_dnn_raw),
                "CHG-EPC-SNSSAI",
            )

        child_snssai_raw = rec.get("__EPC_CHILD_SNSSAI__")
        if child_snssai_raw is not None and clean_cell_value(child_snssai_raw) is not None:
            child_snssai = format_tl1_value(
                "SNSSAI",
                as_text(child_snssai_raw),
                "CHG-EPC-SNSSAI",
            )

        if epc_id not in created_ids:
            create_vals = {
                k: vals[k]
                for k in ("ID", "NAME", "SNSSAI")
                if k in vals
            }

            line = render_cli_line(
                "CRTE-EPC-SNSSAI",
                ["ID", "NAME", "SNSSAI"],
                create_vals,
                required_order=["ID", "NAME", "SNSSAI"],
            )
            if line:
                lines.append(line)
                created_ids.add(epc_id)

            if child_dnn is not None or child_snssai is not None:
                chg_vals = {"ID": epc_id, "MODE": "ADD"}
                if child_dnn is not None:
                    chg_vals["DNN"] = child_dnn
                if child_snssai is not None:
                    chg_vals["SNSSAI"] = child_snssai

                # CHG-EPC-SNSSAI ordering requirement:
                # MODE=ADD must be emitted before DNN, and any child
                # parameters that follow DNN must remain after it.
                chg_order = ["ID", "NAME", "MODE", "DNN", "SNSSAI"]
                line = render_cli_line(
                    "CHG-EPC-SNSSAI",
                    chg_order,
                    chg_vals,
                    required_order=chg_order,
                )
                if line:
                    lines.append(line)
            continue

        if child_dnn is None and child_snssai is None:
            continue

        chg_vals = {"ID": epc_id, "MODE": "ADD"}
        if child_dnn is not None:
            chg_vals["DNN"] = child_dnn
        if child_snssai is not None:
            chg_vals["SNSSAI"] = child_snssai

        # Keep MODE=ADD immediately before DNN for continuation rows too.
        chg_order = ["ID", "NAME", "MODE", "DNN", "SNSSAI"]
        line = render_cli_line(
            "CHG-EPC-SNSSAI",
            chg_order,
            chg_vals,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_snssai_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render SNSSAI parent profile plus repeated DNN mappings correctly."""
    if section.command.upper() != "CRTE-SNSSAI":
        return []

    lines: List[str] = []
    created_ids: set[str] = set()

    create_order = ["ID", "SST", "SD", "EBIALLOC", "NASCGST", "PRI", "DNN", "PDUCONF", "INH"]
    # CHG-SNSSAI-INF ordering requirement:
    # MODE=ADD must appear immediately before DNN, and all child parameters
    # following DNN must remain after it.
    chg_order = ["ID", "DFT", "EBIALLOC", "NASCGST", "PRI", "MODE", "DNN", "PDUCONF", "INH"]

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals or "ID" not in vals or "DNN" not in vals:
            continue

        snssai_id = vals["ID"]

        if snssai_id not in created_ids:
            create_vals = {
                k: vals[k]
                for k in create_order
                if k in vals
            }
            line = render_cli_line(
                "CRTE-SNSSAI",
                create_order,
                create_vals,
                required_order=create_order,
            )
            if line:
                lines.append(line)
                created_ids.add(snssai_id)
            continue

        chg_vals: Dict[str, str] = {
            "ID": snssai_id,
            "MODE": "ADD",
        }
        for k in ("DNN", "PDUCONF", "INH"):
            if k in vals:
                chg_vals[k] = vals[k]

        line = render_cli_line(
            "CHG-SNSSAI-INF",
            chg_order,
            chg_vals,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_ni_pf_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render NI-PF so each ID is created once, then continuation mappings are added.

    Example:
      CRTE-NI-PF:ID=1,NAME=NI-UPF-SPGW,INTF_TYPE=N3N9,NI=NI-UPF-N3N9;
      CHG-NI-PF:ID=1,MODE=ADD,INTF_TYPE=S1U,NI=NI-SGW-S1U;
      CHG-NI-PF:ID=1,MODE=ADD,INTF_TYPE=S5S8U,NI=NI-PGW-S5S8U;
    """
    if section.command.upper() != "CRTE-NI-PF":
        return []

    lines: List[str] = []
    current_id: Optional[str] = None
    current_name: Optional[str] = None
    created_ids: set[str] = set()

    for rec in records:
        vals = record_to_cli_values(section, rec, keep_arrows=False)
        if not vals:
            continue

        if "ID" in vals:
            current_id = vals["ID"]
        if "NAME" in vals:
            current_name = vals["NAME"]

        if current_id is None:
            continue

        intf = vals.get("INTF_TYPE")
        ni = vals.get("NI")
        if not intf or not ni:
            continue

        # First mapping for an ID creates the profile.
        if current_id not in created_ids:
            create_vals = {
                "ID": current_id,
                "INTF_TYPE": intf,
                "NI": ni,
            }
            if current_name:
                create_vals["NAME"] = current_name

            line = render_cli_line(
                "CRTE-NI-PF",
                ["ID", "NAME", "INTF_TYPE", "NI"],
                create_vals,
                required_order=["ID", "NAME", "INTF_TYPE", "NI"],
            )
            if line:
                lines.append(line)
                created_ids.add(current_id)
            continue

        # Additional mappings for the same ID are appended with MODE=ADD.
        chg_vals = {
            "ID": current_id,
            "MODE": "ADD",
            "INTF_TYPE": intf,
            "NI": ni,
        }
        line = render_cli_line(
            "CHG-NI-PF",
            ["ID", "MODE", "INTF_TYPE", "NI"],
            chg_vals,
            required_order=["ID", "MODE", "INTF_TYPE", "NI"],
        )
        if line:
            lines.append(line)

    return lines


def records_to_ne_grp_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render NE-GRP with the first child member on CRTE and later members on CHG.

    Example:
      CRTE-NE-GRP:NE=PCRF,GID=1,GNAME=...,SELM=PRI,TMR=1,MAXCHG=1,NEID=1,PRI=1;
      CHG-NE-GRP:NE=PCRF,GID=1,MODE=ADD,NEID=2,PRI=2;
    """
    if section.command.upper() != "CRTE-NE-GRP":
        return []

    parent_fields = [
        "NE", "GID", "GNAME", "BKUPGID",
        "DFTGID", "DFLTGID", "DEFAULTGID", "DEFGID",
        "SELM", "TMR", "MAXCHG", "AF", "IFDIR"
    ]
    child_fields = ["NEID", "PRI", "CAPA", "ROAM", "RCAPA", "ECR", "INTV"]

    crte_order = [x for x in section.cli_order if x in parent_fields + child_fields]
    if not crte_order:
        crte_order = [
            "NE", "GID", "GNAME", "BKUPGID", "SELM", "TMR", "MAXCHG", "AF",
            "NEID", "PRI", "CAPA", "ROAM", "RCAPA", "ECR", "INTV"
        ]

    chg_order = ["NE", "GID", "MODE", "NEID", "PRI", "CAPA", "ROAM", "RCAPA", "ECR", "INTV"]

    lines: List[str] = []
    current_parent: Dict[str, str] = {}
    pending_parent: Optional[Dict[str, str]] = None
    first_child_emitted = False

    for rec in records:
        values = record_to_cli_values(section, rec, keep_arrows=False)
        if not values:
            continue

        parent_values = {k: v for k, v in values.items() if k in parent_fields}
        child_values = {k: v for k, v in values.items() if k in child_fields}

        # A row with parent fields starts a new group context. Do not emit CRTE
        # yet; wait for the first child so it can be included on the create line.
        if parent_values:
            current_parent = dict(parent_values)
            pending_parent = dict(parent_values)
            first_child_emitted = False

            # If the same record already contains child fields, emit immediately.
            if child_values:
                line_values = {**current_parent, **child_values}
                line = render_cli_line("CRTE-NE-GRP", crte_order, line_values)
                if line:
                    lines.append(line)
                    first_child_emitted = True
                    pending_parent = None
            continue

        if not child_values:
            continue

        # First child for this group belongs on CRTE.
        if not first_child_emitted:
            line_values = {**current_parent, **child_values}
            line = render_cli_line("CRTE-NE-GRP", crte_order, line_values)
            if line:
                lines.append(line)
                first_child_emitted = True
                pending_parent = None
            continue

        # Remaining children are added with MODE=ADD.
        line_values = {k: v for k, v in current_parent.items() if k in {"NE", "GID"}}
        line_values["MODE"] = "ADD"
        line_values.update(child_values)

        line = render_cli_line(
            "CHG-NE-GRP",
            chg_order,
            line_values,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    # Defensive fallback: if a group has no child rows at all, still emit CRTE.
    if pending_parent and not first_child_emitted:
        line = render_cli_line("CRTE-NE-GRP", crte_order, pending_parent)
        if line:
            lines.append(line)

    return lines


def _has_arrow_marker(record: Dict[str, Any]) -> bool:
    return any(is_arrow_value(v) for v in record.values())


def _service_fields_present(section: Section) -> bool:
    order = {x.upper() for x in section.cli_order}
    return "SVC" in order and ("SVCIP" in order or "SVCPORT" in order or "SVCID" in order)


def records_to_peer_service_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Render peer service/IP records as one CRTE followed by CHG additions."""
    cmd = section.command.upper()
    if not cmd.startswith("CRTE-") or cmd == "CRTE-NE-GRP" or not _service_fields_present(section):
        return []

    service_fields = {"SVC", "SCHEME", "SVCIP", "SVCPORT", "SVCID"}
    object_name = cmd[5:]
    chg_command = f"CHG-{object_name}-INF"

    object_key_candidates = {object_name, "ID"}
    parent_key_order = [
        x.upper() for x in section.cli_order
        if x.upper() in object_key_candidates
    ]
    if not parent_key_order and section.cli_order:
        parent_key_order = [section.cli_order[0].upper()]

    chg_order = parent_key_order + [
        "MODE", "SVC", "SCHEME", "SVCIP", "SVCPORT", "SVCID"
    ]

    lines: List[str] = []
    current_parent: Dict[str, str] = {}
    previous_svc: Optional[str] = None
    created_object_key: Optional[Tuple[str, ...]] = None

    for rec in records:
        values = record_to_cli_values(section, rec, keep_arrows=False)
        if not values:
            continue

        parent_values = {
            k: v for k, v in values.items()
            if k.upper() not in service_fields
        }
        svc_values = {
            k: v for k, v in values.items()
            if k.upper() in service_fields
        }

        # Parent fields are repeated by the dedicated extractor so each record
        # remains self-contained.
        current_parent.update(parent_values)

        key_tuple = tuple(
            current_parent.get(k, "")
            for k in parent_key_order
        )

        # First child for a peer object is included directly on CRTE.
        if created_object_key != key_tuple:
            line = render_cli_line(
                cmd,
                section.cli_order,
                {**current_parent, **svc_values},
            )
            if line:
                lines.append(line)
                created_object_key = key_tuple
            previous_svc = svc_values.get("SVC")
            continue

        # Later child rows are additions.
        line_values: Dict[str, str] = {}
        for k in parent_key_order:
            if k in current_parent:
                line_values[k] = current_parent[k]
        line_values["MODE"] = "ADD"
        line_values.update(svc_values)

        current_svc = svc_values.get("SVC")
        # Same service, second+ IP: omit repeated SCHEME.
        if current_svc and previous_svc == current_svc:
            line_values.pop("SCHEME", None)

        previous_svc = current_svc or previous_svc

        line = render_cli_line(
            chg_command,
            chg_order,
            line_values,
            required_order=chg_order,
        )
        if line:
            lines.append(line)

    return lines


def records_to_mmertal_tl1(section: Section, records: List[Dict[str, Any]]) -> List[str]:
    """Special renderer for MMERTAL restricted TA list tables.

    Required behavior:
      - First row with numeric RTALIDX generates:
            CRTE-MMERTAL:RTALIDX=...,CAUSE=...,MCC=...,MNC=...,TAC=...;
      - Following rows after that index generate:
            CHG-MMERTAL-INF:RTALIDX=...,CAUSE=...,TYPE=ADD,MCC=...,MNC=...,TAC=...;
      - Arrow/ditto RTALIDX values inherit the current numeric RTALIDX.
      - TYPE=ADD is inserted only on CHG-MMERTAL-INF rows.
    """
    if section.command.upper() != "CRTE-MMERTAL":
        return []

    lines: List[str] = []
    current_idx: Optional[str] = None
    current_cause: Optional[str] = None
    created_for_idx: set[str] = set()

    crte_order = ["RTALIDX", "CAUSE", "MCC", "MNC", "TAC"]
    chg_order = ["RTALIDX", "CAUSE", "TYPE", "MCC", "MNC", "TAC"]

    for rec in records:
        values = record_to_cli_values(section, rec, keep_arrows=False)
        raw_values = {}
        for param, raw_value in rec.items():
            cli_name = section.param_to_cli.get(param)
            if cli_name:
                raw_values[cli_name.upper()] = raw_value

        # Detect explicit numeric RTALIDX before arrow removal.
        raw_idx = raw_values.get("RTALIDX")
        cleaned_raw_idx = clean_cell_value(raw_idx) if raw_idx is not None and not is_arrow_value(raw_idx) else None
        has_new_idx = cleaned_raw_idx not in (None, "")

        if has_new_idx:
            if "RTALIDX" in values:
                current_idx = values["RTALIDX"]
            else:
                current_idx = format_tl1_value("RTALIDX", cleaned_raw_idx, section.command)

        if "CAUSE" in values and values["CAUSE"] != "":
            current_cause = values["CAUSE"]

        # Skip rows until a numeric RTALIDX has appeared.
        if not current_idx:
            continue

        line_values: Dict[str, str] = {}
        line_values["RTALIDX"] = current_idx

        # CAUSE should be carried forward if blank/arrow after the first row.
        if "CAUSE" in values and values["CAUSE"] != "":
            line_values["CAUSE"] = values["CAUSE"]
        elif current_cause:
            line_values["CAUSE"] = current_cause

        for field in ("MCC", "MNC", "TAC"):
            if field in values:
                line_values[field] = values[field]

        # Do not generate a line without actual TA content unless the first row
        # has all required values in the CRTE line.
        has_ta_content = any(field in line_values and line_values[field] != "" for field in ("MCC", "MNC", "TAC"))
        if not has_ta_content:
            continue

        if current_idx not in created_for_idx:
            line = render_cli_line("CRTE-MMERTAL", crte_order, line_values, required_order=crte_order)
            if line:
                lines.append(line)
                created_for_idx.add(current_idx)
        else:
            line_values["TYPE"] = "ADD"
            line = render_cli_line("CHG-MMERTAL-INF", chg_order, line_values, required_order=chg_order)
            if line:
                lines.append(line)

    return lines



def records_to_conditioned_tl1(
    section: Section,
    records: List[Dict[str, Any]],
) -> List[str]:
    """Render rows with parameter-specific fixed CLI assignments.

    Generic example:
        QCI is unconditional
        DSCP parameter A has DIR=UL
        DSCP parameter B has DIR=DL

    One target row becomes two TL1 lines, one per fixed-condition group.
    """
    conditions = section.param_fixed_cli_values or {}
    if not conditions:
        return []

    lines: List[str] = []
    for record in records:
        if record_has_only_default_non_key_values(section, record):
            continue

        base_values: Dict[str, str] = {}
        grouped: Dict[Tuple[Tuple[str, str], ...], Dict[str, str]] = {}

        for param, raw_value in record.items():
            if str(param).startswith("__"):
                continue

            cli_name = section.param_to_cli.get(param)
            if not cli_name:
                continue

            if is_arrow_value(raw_value):
                continue

            fixed = conditions.get(param, {})

            # A conditioned parameter row represents its own configuration
            # instance. If its target is [default = ...], skip that instance
            # even when other conditioned rows in the same embedded table have
            # explicit values.
            if fixed and is_default_wrapped_value(raw_value):
                continue

            cleaned = clean_cell_value(raw_value)
            if cleaned is None:
                continue

            formatted = format_tl1_value(cli_name, cleaned, section.command)
            if fixed:
                sig = tuple(sorted((k.upper(), v) for k, v in fixed.items()))
                bucket = grouped.setdefault(sig, {})
                bucket[cli_name.upper()] = formatted
            else:
                base_values[cli_name.upper()] = formatted

        # Fixed CLI values from the command template are unconditional.
        base_values.update(extract_fixed_cli_values(section.raw_cli))

        if not grouped:
            continue

        for sig, conditional_values in grouped.items():
            values = dict(base_values)
            for k, v in sig:
                values[k] = format_tl1_value(k, v, section.command)
            values.update(conditional_values)
            line = render_cli_line(section.command, section.cli_order, values)
            if line:
                lines.append(line)

    return lines



def records_to_gtpc_cause_tl1(
    section: Section,
    records: List[Dict[str, Any]],
) -> List[str]:
    """Render hierarchical GTPC-CAUSE child-list rows."""
    if section.command.upper() != "CRTE-GTPC-CAUSE":
        return []

    aliases = {
        "gtpc-cause-index": "INDEX",
        "gtpc-cause-name": "NAME",
        "gtpc-interface-entry/gtpc-cause-interface": "INTF",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-protocol": "PROTC",
        "gtpc-interface-entry/gtpc-protocol-entry/gtpc-cause-error-code": "GTPCC",
    }

    lines: List[str] = []
    current_index = None
    current_name = None
    current_intf = None
    pending_protc = None
    created = False

    for rec in records:
        normalized: Dict[str, Any] = {}
        for param, value in rec.items():
            if str(param).startswith("__"):
                continue
            cli = aliases.get(param) or section.param_to_cli.get(param)
            if cli:
                normalized[cli] = value

        if "INDEX" in normalized:
            current_index = normalized["INDEX"]
        if "NAME" in normalized:
            current_name = normalized["NAME"]

        # Emit CRTE once enough table identity is known.
        if not created and current_index is not None and current_name is not None:
            vals: Dict[str, str] = {}
            ci = clean_cell_value(current_index)
            cn = clean_cell_value(current_name)
            if ci is not None:
                vals["INDEX"] = format_tl1_value("INDEX", ci, "CRTE-GTPC-CAUSE")
            if cn is not None:
                vals["NAME"] = format_tl1_value("NAME", cn, "CRTE-GTPC-CAUSE")
            line = render_cli_line("CRTE-GTPC-CAUSE", ["INDEX", "NAME"], vals)
            if line:
                lines.append(line)
                created = True

        if "INTF" in normalized:
            current_intf = normalized["INTF"]
            pending_protc = None
            continue

        if "PROTC" in normalized:
            pending_protc = normalized["PROTC"]
            continue

        if "GTPCC" in normalized:
            vals: Dict[str, str] = {}
            if current_index is not None:
                v = clean_cell_value(current_index)
                if v is not None:
                    vals["INDEX"] = format_tl1_value("INDEX", v, "CHG-GTPC-CAUSE")
            if current_name is not None:
                v = clean_cell_value(current_name)
                if v is not None:
                    vals["NAME"] = format_tl1_value("NAME", v, "CHG-GTPC-CAUSE")

            vals["MODE"] = "ADD"

            if current_intf is not None:
                v = clean_cell_value(current_intf)
                if v is not None:
                    vals["INTF"] = format_tl1_value("INTF", v, "CHG-GTPC-CAUSE")
            if pending_protc is not None:
                v = clean_cell_value(pending_protc)
                if v is not None:
                    vals["PROTC"] = format_tl1_value("PROTC", v, "CHG-GTPC-CAUSE")

            v = clean_cell_value(normalized["GTPCC"])
            if v is not None:
                vals["GTPCC"] = format_tl1_value("GTPCC", v, "CHG-GTPC-CAUSE")

            line = render_cli_line(
                "CHG-GTPC-CAUSE",
                ["INDEX", "NAME", "MODE", "INTF", "PROTC", "GTPCC"],
                vals,
            )
            if line:
                lines.append(line)

    return lines


def record_to_tl1(section: Section, record: Dict[str, Any]) -> Optional[str]:
    """Convert one target record into one TL1 command line.

    If every assigned non-key value is a CIQ [default = ...] wrapper, skip the
    configuration entirely. Explicit key values may still be present.
    """
    if record_has_only_default_non_key_values(section, record):
        return None

    cli_values = record_to_cli_values(section, record, keep_arrows=False)
    row_action = normalize_row_action(record.get("__ROW_ACTION__"))
    command = derive_command_for_row(section.command, row_action)
    command_upper = command.upper()

    # Skip Service Overload Control entries that target the synthetic
    # "{all services}" selector.
    if command_upper == "CHG-OVLD-CNTL":
        svc = cli_values.get("SVC", "")
        if svc.strip().strip('"').lower() == "{all services}":
            return None

    # Skip VES control entries when EVENT selects all events.
    # Accept both source variants:
    #   EVENT=ALL
    #   EVENT={ALL}
    if command_upper == "CHG-VES-CNTL":
        event = cli_values.get("EVENT", "")
        normalized_event = event.strip().strip('"').strip("{}").strip().upper()
        if normalized_event == "ALL":
            return None

    # MEAS-KDAY TL1 expects TYPE without the RAW- prefix.
    # RAW-1MIN -> 1MIN, RAW-5MIN -> 5MIN, RAW-HOUR -> HOUR, etc.
    if command_upper == "CHG-MEAS-KDAY" and "TYPE" in cli_values:
        type_value = cli_values["TYPE"]
        if type_value.upper().startswith("RAW-"):
            cli_values["TYPE"] = type_value[4:]

    return render_cli_line(command, section.cli_order, cli_values)


def build_tl1_output(wb_values, selected_sheets: Sequence[str], selected_section_numbers: Optional[set[int]] = None) -> Tuple[str, List[str]]:
    sections = discover_sections(wb_values, selected_sheets)
    output_blocks: List[str] = []
    warnings: List[str] = []

    for idx, section in enumerate(sections, start=1):
        if selected_section_numbers and idx not in selected_section_numbers:
            continue

        ws_for_size = wb_values[section.sheet_name]

        next_cli_row = min(
            [s.cli_row for s in sections
             if s.sheet_name == section.sheet_name and s.cli_row > section.cli_row]
            or [ws_for_size.max_row + 1]
        )

        # Skip the entire section before parsing any target values when either
        # Review Status or Config Status says Out of Scope.
        if is_out_of_scope_section(ws_for_size, section, next_cli_row):
            warnings.append(
                f"Skipped Out of Scope section #{idx}: {section.title} "
                f"[{section.sheet_name} row {section.cli_row}]"
            )
            continue

        # MEAS-KDAY is intentionally excluded from TL1 generation.
        if section.command.upper() == "CHG-MEAS-KDAY" or "MEAS-KDAY" in section.title.upper():
            warnings.append(
                f"Skipped MEAS-KDAY section #{idx}: {section.title} "
                f"[{section.sheet_name} row {section.cli_row}]"
            )
            continue

        if getattr(ws_for_size, "max_row", 0) and ws_for_size.max_row >= LARGE_SHEET_ROW_THRESHOLD:
            print(f"Processing large sheet '{section.sheet_name}' with {ws_for_size.max_row:,} rows, section #{idx}: {section.title}")

        extracted_records = extract_target_records_for_section(wb_values, section, sections)

        if section_has_unmapped_pdu_dscp_children(wb_values, section, sections):
            warnings.append(
                f"PDU-CONF section #{idx} contains DSCP child-list Target Values "
                f"(FQI/UL/DL) but the CIQ source does not provide a TL1 CLI mapping "
                f"for those child fields; parent CRTE-PDU-CONF is generated, child "
                f"DSCP rows are not guessed. [{section.sheet_name} row {section.cli_row}]"
            )

        # Generic default-only filtering must also apply to special renderers
        # (NE-GRP, peer service entries, MMERTAL, ACT-LIC, etc.) that do not
        # route through record_to_tl1().
        # Generic default-only filtering applies to every section, including
        # CHG-SUPP-LMT. A key plus only [default = ...] non-key values does not
        # generate a TL1 command.
        default_only_count = sum(
            1 for rec in extracted_records
            if record_has_only_default_non_key_values(section, rec)
        )
        records = [
            rec for rec in extracted_records
            if not record_has_only_default_non_key_values(section, rec)
        ]

        if len(records) >= 50000:
            print(f"  Rendering {len(records):,} TL1 candidate records for section #{idx}: {section.title}")

        if section.command.upper() == "ACT-LIC":
            lines = records_to_act_lic_tl1(section, records)
        elif section.command.upper() == "CHG-OFC-EVT":
            lines = records_to_ofc_evt_tl1(section, records)
        elif section.command.upper() == "CHG-ONC-RCODE":
            lines = records_to_onc_rcode_tl1(section, records)
        elif section.command.upper() == "CRTE-RBASE":
            lines = records_to_rbase_tl1(section, records)
        elif section.command.upper() == "CRTE-PDU-CONF":
            lines = records_to_pdu_conf_tl1(section, records)
        elif section.command.upper() == "CRTE-PCF-PF":
            lines = records_to_pcf_pf_tl1(section, records)
        elif section.command.upper() == "CRTE-EPC-SNSSAI":
            lines = records_to_epc_snssai_tl1(section, records)
        elif section.command.upper() == "CRTE-SNSSAI":
            lines = records_to_snssai_tl1(section, records)
        elif section.command.upper() == "CRTE-NI-PF":
            lines = records_to_ni_pf_tl1(section, records)
        elif section.command.upper() == "CRTE-GTPC-CAUSE":
            lines = records_to_gtpc_cause_tl1(section, records)
        elif section.param_fixed_cli_values:
            lines = records_to_conditioned_tl1(section, records)
        elif section.command.upper() == "CRTE-NE-GRP":
            lines = records_to_ne_grp_tl1(section, records)
        elif section.command.upper() == "CRTE-MMERTAL":
            lines = records_to_mmertal_tl1(section, records)
        else:
            lines = records_to_peer_service_tl1(section, records)
            if not lines:
                lines = []
                for row_num, rec in enumerate(records, start=1):
                    line = record_to_tl1(section, rec)
                    if line:
                        lines.append(line)
                    if row_num % 50000 == 0:
                        print(f"  Rendered {row_num:,} records for section #{idx}...")

        if lines:
            block = [f"### {section.title}"] + lines
            output_blocks.append("\n".join(block))
            if default_only_count:
                warnings.append(
                    f"Skipped {default_only_count} default-only configuration row(s) "
                    f"for section #{idx}: {section.title} "
                    f"[{section.sheet_name} row {section.cli_row}]"
                )
        else:
            if extracted_records and default_only_count == len(extracted_records):
                warnings.append(
                    f"Skipped default-only section #{idx}: {section.title} "
                    f"({default_only_count} row(s)) "
                    f"[{section.sheet_name} row {section.cli_row}]"
                )
            else:
                warnings.append(
                    f"No TL1 rows generated for section #{idx}: {section.title} "
                    f"[{section.sheet_name} row {section.cli_row}]"
                )

    return "\n\n".join(output_blocks).rstrip() + "\n", warnings


# -----------------------------
# User interface
# -----------------------------
def choose_file_dialog() -> Optional[str]:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        path = filedialog.askopenfilename(
            title="Select Excel workbook",
            filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        root.destroy()
        return path or None
    except Exception:
        return None


def choose_folder_dialog() -> Optional[str]:
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        path = filedialog.askdirectory(title="Select output folder")
        root.destroy()
        return path or None
    except Exception:
        return None


def prompt_path(prompt: str, must_exist: bool = False, is_dir: bool = False) -> str:
    while True:
        p = input(prompt).strip().strip('"')
        if not p:
            continue
        path = Path(p)
        if must_exist:
            if is_dir and not path.is_dir():
                print("Folder not found. Try again.")
                continue
            if not is_dir and not path.is_file():
                print("File not found. Try again.")
                continue
        return str(path)


def ask_workbook_path() -> str:
    print("Select Excel workbook...")
    path = choose_file_dialog()
    if path:
        return path
    return prompt_path("Enter Excel workbook path: ", must_exist=True)




def split_text_by_max_lines(text: str, max_lines: int = MAX_TL1_LINES_PER_FILE) -> List[str]:
    """Split output text into chunks with at most max_lines lines per file.

    Blank separator lines and section headers count as lines, so every generated
    file is guaranteed to stay within the configured maximum line count.
    """
    lines = text.splitlines()
    if not lines:
        return [""]

    chunks: List[str] = []
    for start in range(0, len(lines), max_lines):
        chunk_lines = lines[start:start + max_lines]
        chunks.append("\n".join(chunk_lines).rstrip() + "\n")
    return chunks



def timestamped_run_output_path(output_path: str) -> Path:
    """Return an output path inside a YYYYMMDD_HHMMSS run folder.

    The caller supplies the desired parent directory and filename. This helper
    guarantees that generation occurs inside a timestamped child folder:

        <selected-output-folder>/YYYYMMDD_HHMMSS/<filename>

    If the supplied parent is already a timestamp run folder, it is reused.
    This keeps CLI and GUI behavior identical and prevents accidental double
    timestamp nesting.
    """
    path = Path(output_path)
    parent = path.parent

    if re.fullmatch(r"\d{8}_\d{6}", parent.name):
        run_folder = parent
    else:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        run_folder = parent / timestamp

        # A second execution within the exact same second should not mix files.
        if run_folder.exists():
            suffix = 2
            base = run_folder
            while run_folder.exists():
                run_folder = parent / f"{base.name}_{suffix:02d}"
                suffix += 1

    run_folder.mkdir(parents=True, exist_ok=True)
    return run_folder / path.name



def write_warnings_file(written_files: Sequence[str], warnings: Sequence[str]) -> Optional[str]:
    """Write warnings beside the actual TL1 output inside the run timestamp folder."""
    if not warnings or not written_files:
        return None

    first_output = Path(written_files[0])
    warn_path = first_output.with_suffix(".warnings.txt")
    warn_path.write_text("\n".join(warnings) + "\n", encoding="utf-8")
    return str(warn_path)


def write_tl1_files(output_path: str, text: str, max_lines: int = MAX_TL1_LINES_PER_FILE) -> List[str]:
    """Write TL1 output, splitting into numbered files when needed.

    If the output has <= max_lines lines, writes the requested output_path.
    If it has more, writes:
        <stem>_part001.tl1
        <stem>_part002.tl1
        ...
    and does not create the oversized unsplit file.
    """
    path = timestamped_run_output_path(output_path)

    chunks = split_text_by_max_lines(text, max_lines=max_lines)
    written: List[str] = []

    if len(chunks) == 1:
        with open(path, "w", encoding="utf-8", newline="\n") as f:
            f.write(chunks[0])
        written.append(str(path))
        return written

    for idx, chunk in enumerate(chunks, start=1):
        part_path = path.with_name(f"{path.stem}_part{idx:03d}{path.suffix}")
        with open(part_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(chunk)
        written.append(str(part_path))

    return written

def ask_output_path(default_stem: str) -> str:
    """Choose the parent output folder and desired TL1 filename.

    The timestamped run directory is created centrally by write_tl1_files(), so
    every caller follows the same output-folder behavior.
    """
    print("Select output folder...")
    folder = choose_folder_dialog()
    if not folder:
        folder = prompt_path("Enter output folder path: ", must_exist=True, is_dir=True)

    name = input(f"Enter output TL1 filename stem [{default_stem}]: ").strip().strip('"')
    if name:
        stem = Path(name).stem if name.lower().endswith(".tl1") else name
    else:
        stem = default_stem

    filename = f"{stem}_{VERSION_TAG}.tl1"
    return str(Path(folder) / filename)

def ask_sheet_selection(wb_values) -> List[str]:
    print("\nWorkbook sheets:")
    for i, name in enumerate(wb_values.sheetnames, start=1):
        print(f"  {i}. {name}")

    choice = input("\nProcess whole workbook or specific sheet(s)? [A=all / S=specific] (default A): ").strip().lower()
    if choice not in {"s", "specific"}:
        large_sheets = [name for name in wb_values.sheetnames if wb_values[name].max_row >= LARGE_SHEET_ROW_THRESHOLD]
        if large_sheets:
            print("\nLarge sheets detected:")
            for name in large_sheets:
                print(f"  - {name}: {wb_values[name].max_row:,} rows")
            confirm = input("Process ALL sheets including these large sheets? [y/N] (default N): ").strip().lower()
            if confirm not in {"y", "yes"}:
                print("Large sheets will be skipped during CLI-section discovery. They can still be read when referenced by a selected section.")
                return [name for name in wb_values.sheetnames if name not in set(large_sheets)]
        return list(wb_values.sheetnames)

    raw = input("Enter sheet number(s) or name(s), comma-separated: ").strip()
    selected = []
    for token in [x.strip() for x in raw.split(",") if x.strip()]:
        range_match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", token)
        if range_match:
            start_idx, end_idx = int(range_match.group(1)), int(range_match.group(2))
            if start_idx > end_idx:
                start_idx, end_idx = end_idx, start_idx
            for idx in range(start_idx, end_idx + 1):
                if 1 <= idx <= len(wb_values.sheetnames):
                    selected.append(wb_values.sheetnames[idx - 1])
                else:
                    print(f"WARNING: sheet number out of range: {idx}")
        elif token.isdigit():
            idx = int(token)
            if 1 <= idx <= len(wb_values.sheetnames):
                selected.append(wb_values.sheetnames[idx - 1])
        elif token in wb_values.sheetnames:
            selected.append(token)
        else:
            print(f"WARNING: sheet not found: {token}")

    return selected or list(wb_values.sheetnames)


def ask_section_selection(sections: List[Section]) -> Optional[set[int]]:
    if not sections:
        return None
    print("\nDiscovered CLI sections:")
    for i, sec in enumerate(sections, start=1):
        print(f"  {i}. [{sec.sheet_name}] {sec.title} -> {sec.command}")
    raw = input("\nProcess all sections or selected section number(s)? [A=all / numbers comma-separated] (default A): ").strip()
    if not raw or raw.lower() in {"a", "all"}:
        return None
    nums = set()
    for token in raw.split(","):
        token = token.strip()
        range_match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", token)
        if range_match:
            start_idx, end_idx = int(range_match.group(1)), int(range_match.group(2))
            if start_idx > end_idx:
                start_idx, end_idx = end_idx, start_idx
            nums.update(range(start_idx, end_idx + 1))
        elif token.isdigit():
            nums.add(int(token))
    return nums or None


def main() -> None:
    print(f"Excel CIQ -> TL1 generator {VERSION_TAG}")
    workbook_path = ask_workbook_path()
    print(f"Loading workbook: {workbook_path}")

    # data_only=True uses cached formula results saved by Excel. This is important because many CIQ templates use formulas like =C12.
    # For very large workbooks, set TL1_READ_ONLY=1 before running the script to reduce memory usage.
    # Normal mode remains the default because this parser uses random cell access heavily.
    read_only_mode = os.environ.get("TL1_READ_ONLY", "").strip().lower() in {"1", "true", "yes", "y"}
    if read_only_mode:
        print("Large-workbook mode enabled: loading workbook with openpyxl read_only=True.")
    wb_values = load_workbook(workbook_path, data_only=True, read_only=read_only_mode)

    selected_sheets = ask_sheet_selection(wb_values)
    sections = discover_sections(wb_values, selected_sheets)
    if not sections:
        print("ERROR: No sections found. Expected cells containing 'CLI Command:'.")
        sys.exit(2)

    selected_section_numbers = ask_section_selection(sections)
    default_stem = Path(workbook_path).stem + "_output"
    output_path = ask_output_path(default_stem)

    text, warnings = build_tl1_output(wb_values, selected_sheets, selected_section_numbers)

    written_files = write_tl1_files(output_path, text, max_lines=MAX_TL1_LINES_PER_FILE)

    print("\nDone. TL1 file generated:")
    for path in written_files:
        print(f"  {path}")
    if len(written_files) > 1:
        print(f"Split output into {len(written_files)} files with maximum 10000 lines per file.")

    line_count = sum(1 for line in text.splitlines() if line.strip() and not line.startswith("###"))
    total_lines = len(text.splitlines())
    print(f"Generated TL1 command lines: {line_count}")
    print(f"Generated total output lines: {total_lines}")

    warn_path = write_warnings_file(written_files, warnings)
    if warn_path:
        print(f"Warnings written to: {warn_path}")


if __name__ == "__main__":
    main()
