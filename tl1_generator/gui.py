from __future__ import annotations

import os
import subprocess
import sys
import threading
import traceback
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter.scrolledtext import ScrolledText

from openpyxl import load_workbook

from . import APP_VERSION, ENGINE_VERSION
from . import engine


class TL1GeneratorApp(tk.Tk):
    """Cross-platform desktop front end for the CIQ -> TL1 generator."""

    def __init__(self) -> None:
        super().__init__()
        self.title(f"Tools TL1 Generator v{APP_VERSION}")
        self.geometry("1120x760")
        self.minsize(980, 680)

        self.workbook_path = tk.StringVar()
        self.output_folder = tk.StringVar(value=str(Path.home()))
        self.output_stem = tk.StringVar(value="ciq_tl1_output")
        self.max_lines = tk.IntVar(value=10000)
        self.status_text = tk.StringVar(value="Ready")

        self.wb = None
        self.sections: List[engine.Section] = []
        self.scanned_sheet_names: List[str] = []
        self.busy = False

        self._build_ui()
        self._build_menu()

    # -----------------------------
    # UI construction
    # -----------------------------
    def _build_menu(self) -> None:
        menubar = tk.Menu(self)
        file_menu = tk.Menu(menubar, tearoff=False)
        file_menu.add_command(label="Open workbook...", command=self.browse_workbook)
        file_menu.add_command(label="Open output folder", command=self.open_output_folder)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.destroy)
        menubar.add_cascade(label="File", menu=file_menu)

        help_menu = tk.Menu(menubar, tearoff=False)
        help_menu.add_command(label="About", command=self.show_about)
        menubar.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menubar)

    def _build_ui(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass

        top = ttk.Frame(self, padding=(16, 14, 16, 8))
        top.pack(fill="x")

        title = ttk.Label(top, text="Tools TL1 Generator", font=("Arial", 20, "bold"))
        title.pack(anchor="w")
        subtitle = ttk.Label(
            top,
            text=f"Desktop app v{APP_VERSION} • TL1 engine v{ENGINE_VERSION} • Windows / macOS / Linux",
        )
        subtitle.pack(anchor="w", pady=(2, 0))

        paths = ttk.LabelFrame(self, text="Files", padding=12)
        paths.pack(fill="x", padx=16, pady=(0, 10))

        ttk.Label(paths, text="Excel workbook").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(paths, textvariable=self.workbook_path).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Button(paths, text="Browse...", command=self.browse_workbook).grid(row=0, column=2, padx=(8, 0), pady=4)
        ttk.Button(paths, text="Load", command=self.load_workbook_async).grid(row=0, column=3, padx=(8, 0), pady=4)

        ttk.Label(paths, text="Output folder").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(paths, textvariable=self.output_folder).grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Button(paths, text="Browse...", command=self.browse_output_folder).grid(row=1, column=2, padx=(8, 0), pady=4)
        ttk.Button(paths, text="Open", command=self.open_output_folder).grid(row=1, column=3, padx=(8, 0), pady=4)

        ttk.Label(paths, text="Output filename stem").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=4)
        ttk.Entry(paths, textvariable=self.output_stem).grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(paths, text="Max lines/file").grid(row=2, column=2, sticky="e", padx=(8, 8), pady=4)
        ttk.Spinbox(paths, from_=1000, to=100000, increment=1000, textvariable=self.max_lines, width=10).grid(
            row=2, column=3, sticky="w", pady=4
        )
        paths.columnconfigure(1, weight=1)

        middle = ttk.Panedwindow(self, orient="horizontal")
        middle.pack(fill="both", expand=True, padx=16, pady=(0, 10))

        sheet_frame = ttk.LabelFrame(middle, text="Sheets", padding=8)
        section_frame = ttk.LabelFrame(middle, text="CLI sections", padding=8)
        middle.add(sheet_frame, weight=1)
        middle.add(section_frame, weight=2)

        self.sheet_list = tk.Listbox(sheet_frame, selectmode=tk.EXTENDED, exportselection=False)
        sheet_scroll = ttk.Scrollbar(sheet_frame, orient="vertical", command=self.sheet_list.yview)
        self.sheet_list.configure(yscrollcommand=sheet_scroll.set)
        self.sheet_list.grid(row=0, column=0, columnspan=3, sticky="nsew")
        sheet_scroll.grid(row=0, column=3, sticky="ns")
        ttk.Button(sheet_frame, text="Select all", command=self.select_all_sheets).grid(row=1, column=0, sticky="ew", pady=(8, 0))
        ttk.Button(sheet_frame, text="Clear", command=lambda: self.sheet_list.selection_clear(0, tk.END)).grid(
            row=1, column=1, sticky="ew", padx=6, pady=(8, 0)
        )
        ttk.Button(sheet_frame, text="Scan sections", command=self.scan_sections_async).grid(row=1, column=2, sticky="ew", pady=(8, 0))
        sheet_frame.rowconfigure(0, weight=1)
        sheet_frame.columnconfigure(0, weight=1)
        sheet_frame.columnconfigure(1, weight=1)
        sheet_frame.columnconfigure(2, weight=1)

        self.section_list = tk.Listbox(section_frame, selectmode=tk.EXTENDED, exportselection=False)
        section_scroll = ttk.Scrollbar(section_frame, orient="vertical", command=self.section_list.yview)
        self.section_list.configure(yscrollcommand=section_scroll.set)
        self.section_list.grid(row=0, column=0, columnspan=3, sticky="nsew")
        section_scroll.grid(row=0, column=3, sticky="ns")
        ttk.Button(section_frame, text="Select all", command=self.select_all_sections).grid(row=1, column=0, sticky="ew", pady=(8, 0))
        ttk.Button(section_frame, text="Clear", command=lambda: self.section_list.selection_clear(0, tk.END)).grid(
            row=1, column=1, sticky="ew", padx=6, pady=(8, 0)
        )
        ttk.Button(section_frame, text="Generate TL1", command=self.generate_async).grid(row=1, column=2, sticky="ew", pady=(8, 0))
        section_frame.rowconfigure(0, weight=1)
        section_frame.columnconfigure(0, weight=1)
        section_frame.columnconfigure(1, weight=1)
        section_frame.columnconfigure(2, weight=1)

        log_frame = ttk.LabelFrame(self, text="Activity", padding=8)
        log_frame.pack(fill="both", padx=16, pady=(0, 10))
        self.log = ScrolledText(log_frame, height=8, wrap="word", state="disabled", font=("Courier New", 9))
        self.log.pack(fill="both", expand=True)

        status_bar = ttk.Frame(self, padding=(16, 4, 16, 10))
        status_bar.pack(fill="x")
        ttk.Label(status_bar, textvariable=self.status_text).pack(side="left")
        self.progress = ttk.Progressbar(status_bar, mode="indeterminate", length=180)
        self.progress.pack(side="right")

    # -----------------------------
    # Utility helpers
    # -----------------------------
    def append_log(self, message: str) -> None:
        def _write() -> None:
            self.log.configure(state="normal")
            self.log.insert(tk.END, message.rstrip() + "\n")
            self.log.see(tk.END)
            self.log.configure(state="disabled")

        self.after(0, _write)

    def set_busy(self, busy: bool, status: str) -> None:
        self.busy = busy

        def _update() -> None:
            self.status_text.set(status)
            if busy:
                self.progress.start(12)
            else:
                self.progress.stop()

        self.after(0, _update)

    def run_worker(self, target, status: str) -> None:
        if self.busy:
            messagebox.showinfo("Busy", "A task is already running.")
            return

        def wrapper() -> None:
            self.set_busy(True, status)
            try:
                target()
            except Exception as exc:  # noqa: BLE001
                self.append_log(f"ERROR: {exc}")
                self.append_log(traceback.format_exc())
                self.after(0, lambda: messagebox.showerror("Error", str(exc)))
            finally:
                self.set_busy(False, "Ready")

        threading.Thread(target=wrapper, daemon=True).start()

    def browse_workbook(self) -> None:
        path = filedialog.askopenfilename(
            title="Select CIQ Excel workbook",
            filetypes=[("Excel workbook", "*.xlsx *.xlsm"), ("All files", "*.*")],
        )
        if path:
            self.workbook_path.set(path)
            self.output_stem.set(Path(path).stem + "_output")

    def browse_output_folder(self) -> None:
        folder = filedialog.askdirectory(title="Select output folder")
        if folder:
            self.output_folder.set(folder)

    def open_output_folder(self) -> None:
        folder = Path(self.output_folder.get().strip() or Path.home())
        folder.mkdir(parents=True, exist_ok=True)
        try:
            if sys.platform.startswith("win"):
                os.startfile(str(folder))  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(folder)])
            else:
                subprocess.Popen(["xdg-open", str(folder)])
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Open folder", str(exc))

    def show_about(self) -> None:
        messagebox.showinfo(
            "About",
            f"Tools TL1 Generator v{APP_VERSION}\n"
            f"Bundled TL1 engine v{ENGINE_VERSION}\n\n"
            "Converts CIQ Excel Target Values into TL1 command files.\n"
            "Supports multi-sheet workbooks, section selection, warnings, and output splitting.",
        )

    # -----------------------------
    # Workbook / section handling
    # -----------------------------
    def load_workbook_async(self) -> None:
        self.run_worker(self._load_workbook, "Loading workbook...")

    def _load_workbook(self) -> None:
        path = Path(self.workbook_path.get().strip())
        if not path.is_file():
            raise FileNotFoundError("Select a valid .xlsx or .xlsm workbook first.")

        self.append_log(f"Loading workbook: {path}")
        wb = load_workbook(path, data_only=True)
        self.wb = wb

        def populate() -> None:
            self.sheet_list.delete(0, tk.END)
            self.section_list.delete(0, tk.END)
            self.sections = []
            self.scanned_sheet_names = []
            for name in wb.sheetnames:
                ws = wb[name]
                suffix = f"  ({ws.max_row:,} rows)" if ws.max_row >= 100000 else ""
                self.sheet_list.insert(tk.END, name + suffix)
            self.select_all_sheets()

        self.after(0, populate)
        self.append_log(f"Loaded {len(wb.sheetnames)} sheet(s).")
        large = [name for name in wb.sheetnames if wb[name].max_row >= 100000]
        if large:
            self.append_log("Large sheet(s) detected: " + ", ".join(large))
            self.append_log("Tip: select only the required configuration sheets before scanning sections.")

    def select_all_sheets(self) -> None:
        self.sheet_list.selection_set(0, tk.END)

    def select_all_sections(self) -> None:
        self.section_list.selection_set(0, tk.END)

    def selected_sheet_names(self) -> List[str]:
        if self.wb is None:
            return []
        selected = list(self.sheet_list.curselection())
        if not selected:
            return list(self.wb.sheetnames)
        return [self.wb.sheetnames[i] for i in selected]

    def scan_sections_async(self) -> None:
        self.run_worker(self._scan_sections, "Scanning CLI sections...")

    def _scan_sections(self) -> None:
        if self.wb is None:
            raise RuntimeError("Load a workbook first.")
        sheets = self.selected_sheet_names()
        self.append_log("Scanning sheets: " + ", ".join(sheets))
        sections = engine.discover_sections(self.wb, sheets)
        self.sections = sections
        self.scanned_sheet_names = sheets

        def populate() -> None:
            self.section_list.delete(0, tk.END)
            for i, sec in enumerate(sections, start=1):
                self.section_list.insert(tk.END, f"{i}. [{sec.sheet_name}] {sec.title} → {sec.command}")
            self.select_all_sections()

        self.after(0, populate)
        self.append_log(f"Found {len(sections)} CLI section(s).")

    # -----------------------------
    # Generation
    # -----------------------------
    def generate_async(self) -> None:
        self.run_worker(self._generate, "Generating TL1...")

    def _generate(self) -> None:
        if self.wb is None:
            raise RuntimeError("Load a workbook first.")

        selected_sheets = self.selected_sheet_names()
        if not selected_sheets:
            raise RuntimeError("Select at least one sheet.")

        # Section numbers are valid only if the current sheet selection matches the last scan.
        selected_section_numbers: Optional[set[int]] = None
        if self.sections and selected_sheets == self.scanned_sheet_names:
            selected_rows = list(self.section_list.curselection())
            if selected_rows and len(selected_rows) < len(self.sections):
                selected_section_numbers = {i + 1 for i in selected_rows}
        elif self.sections:
            self.append_log("Sheet selection changed since the last section scan; processing all sections in selected sheets.")

        output_folder = Path(self.output_folder.get().strip())
        output_folder.mkdir(parents=True, exist_ok=True)
        stem = self.output_stem.get().strip() or "ciq_tl1_output"
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_folder / f"{stem}_{engine.VERSION_TAG}_{timestamp}.tl1"

        max_lines = int(self.max_lines.get() or 10000)
        if max_lines < 1:
            raise ValueError("Max lines/file must be at least 1.")

        self.append_log(f"Generating from {len(selected_sheets)} selected sheet(s)...")
        text, warnings = engine.build_tl1_output(self.wb, selected_sheets, selected_section_numbers)
        written = engine.write_tl1_files(str(output_path), text, max_lines=max_lines)

        for file_path in written:
            self.append_log(f"Generated: {file_path}")

        if warnings:
            warn_path = output_path.with_suffix(".warnings.txt")
            warn_path.write_text("\n".join(warnings) + "\n", encoding="utf-8")
            self.append_log(f"Warnings: {warn_path}")
            self.append_log(f"Warning count: {len(warnings)}")
        else:
            self.append_log("No warnings generated.")

        command_count = sum(1 for line in text.splitlines() if line.strip() and not line.startswith("###"))
        self.append_log(f"TL1 command lines: {command_count:,}")
        self.after(
            0,
            lambda: messagebox.showinfo(
                "Generation complete",
                f"Generated {command_count:,} TL1 command line(s)\n"
                f"Output file(s): {len(written)}\n"
                f"Warnings: {len(warnings)}",
            ),
        )


def main() -> None:
    app = TL1GeneratorApp()
    app.mainloop()


if __name__ == "__main__":
    main()
